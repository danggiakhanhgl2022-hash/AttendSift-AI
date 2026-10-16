"""Excel presentation and imported Chatwork heading regressions (14 Oct 2026)."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

import openpyxl

from src.services.event_policy import is_annual_leave_text_only, mark_annual_leave_source_only
from src.utils.excel_export import export_to_excel
from src.utils.message_cleaning import deduplicate_copied_sender_header


def _message(mid, text, name='Lê Đức Quân', date='2026-10-09'):
    return {'message_id': mid, 'room_id': 'room', 'account_id': 7,
            'name': name, 'body': text, 'send_time': date + 'T08:13:26+07:00'}


def _event(mid, name='Quân'):
    return {'event_id': 'EVT00001', 'person_id': name, 'employee_name': name,
            'event_date': '2026-10-08', 'event_type': 'Tăng ca',
            'ot_minutes': 240, 'ot_used_minutes': 0,
            'source_message_ids': [mid], 'is_excluded': False}


def test_repeated_copied_chatwork_sender_headers_only_once():
    raw = ('クアン (Lê Đức Quân)\n'
           'クアン (Lê Đức Quân)B PRODUCTIONS COMPANY LIMITED\n'
           'Team DemoBeta: ngày 8/10\nĐăng, Bảo, Quân, Việt OT 4h')
    assert deduplicate_copied_sender_header(raw, 'Lê Đức Quân') == (
        'クアン (Lê Đức Quân)\n'
        'Team DemoBeta: ngày 8/10\nĐăng, Bảo, Quân, Việt OT 4h')
    # No change to ordinary content where a person's name is mentioned again.
    plain = 'Team DemoBeta: ngày 8/10\nQuân OT 4h\nQuân hoàn thành công việc.'
    assert deduplicate_copied_sender_header(plain, 'Lê Đức Quân') == plain


def test_export_raw_cells_remove_only_duplicate_header_and_keep_original_source(tmp_path):
    raw = ('クアン (Lê Đức Quân)\n'
           'クアン (Lê Đức Quân)B PRODUCTIONS COMPANY LIMITED\n'
           'Team DemoBeta: ngày 8/10\nĐăng, Bảo, Quân, Việt OT 4h\n')
    second = 'Thư OT 1h30\nAnh Tâm OT 30ph'
    snapshot = [_message('m1', raw), _message('m2', second, name='Người báo cáo khác')]
    path = tmp_path / 'cham_cong.xlsx'
    export_to_excel([_event('m1')], {'month': 10, 'year': 2026,
                                    'snapshot': snapshot}, path)
    assert snapshot[0]['body'] == raw  # never alter evidence in the DB/snapshot
    wb = openpyxl.load_workbook(path)
    try:
        raw_expected = 'クアン (Lê Đức Quân)\nTeam DemoBeta: ngày 8/10\nĐăng, Bảo, Quân, Việt OT 4h\n'
        assert wb['1_TIN_NHAN_GOC']['E5'].value == raw_expected
        assert wb['2_CHAM_CONG']['Z5'].value == raw_expected
        assert wb['1_TIN_NHAN_GOC']['E6'].value == second
        assert wb['1_TIN_NHAN_GOC']['E5'].alignment.wrap_text
        assert wb['1_TIN_NHAN_GOC'].row_dimensions[5].height > 32
        for sheet_name in ['1_TIN_NHAN_GOC', '2_CHAM_CONG']:
            tab = wb[sheet_name]
            assert tab['A5'].fill.fgColor.rgb[-6:] == 'CDEEF8'
            assert tab['A5'].border.bottom.style == 'thin'
        assert wb['1_TIN_NHAN_GOC']['A6'].fill.fgColor.rgb[-6:] == 'FFFFFF'
        assert wb['1_TIN_NHAN_GOC'].tables['TinNhanGocTable'].ref == 'A4:H6'
        assert wb['2_CHAM_CONG'].tables['ChamCongTable'].ref == 'A4:Z5'
        assert wb['3_TONG_KET'].tables['TongKetTable'].ref == 'A4:M5'
        assert wb['4_CAN_KIEM_TRA'].tables['CanKiemTraTable'].ref == 'A4:G5'
    finally:
        wb.close()
    # ZIP-level integrity: tables' autoFilter ranges and table ranges agree.
    with ZipFile(path) as z:
        for name in [n for n in z.namelist() if n.startswith('xl/tables/') and n.endswith('.xml')]:
            root = ElementTree.fromstring(z.read(name))
            ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
            filt = root.find(f'{ns}autoFilter')
            assert filt is None or filt.attrib.get('ref') == root.attrib['ref']


def test_export_longer_than_template_range_still_extends_excel_table(tmp_path):
    # The prior template had a fixed range A4:H193; new messages must get stripes.
    snapshot = [_message(f'm{i}',f'Tin thật {i}') for i in range(200)]
    path = tmp_path / 'nhieu_tin.xlsx'
    export_to_excel([], {'month':10,'year':2026,'snapshot':snapshot},path)
    wb=openpyxl.load_workbook(path,read_only=False)
    try:
        sh=wb['1_TIN_NHAN_GOC']
        assert sh.tables['TinNhanGocTable'].ref == 'A4:H204'
        assert sh['A204'].fill.fgColor.rgb[-6:] == 'FFFFFF'
        assert sh['E204'].value == 'Tin thật 199'
    finally:
        wb.close()


def test_annual_leave_off_half_day_is_source_only_and_not_a_review_warning():
    text=('Em xin phép off sáng T6 ngày 16/10 (0.5 phép năm )\n'
          'Lí do: Em lên trường làm giấy GDQP và NVQS\nEm xin cảm ơn')
    msg = _message('m-leave', text, name='Đỗ Gia Quân')
    event = {'source_message_ids':['m-leave'],'event_type':'Nghỉ',
             'leave_days':None,'needs_review':True,'validation_errors':['leave_days: cần xác định'],
             'validation_warnings':['leave_days: cần xác định']}
    assert is_annual_leave_text_only(text)
    assert mark_annual_leave_source_only(event,[msg])
    assert event['source_only'] and event['is_excluded']
    assert not event['needs_review']
    assert event['validation_errors'] == []


def test_half_day_annual_leave_off_exports_information_not_attendance(tmp_path):
    body = 'Em xin phép off sáng T6 ngày 16/10 (0.5 phép năm)\nEm xin cảm ơn.'
    msg = _message('leave', body, name='Đỗ Gia Quân')
    event = {'event_id':'EVT100','person_id':'P', 'employee_name':'Đỗ Gia Quân',
             'event_date':'2026-10-16', 'event_type':'Nghỉ', 'event_subtype':'Phép năm',
             'leave_days':None, 'source_message_ids':['leave'], 'needs_review':True}
    assert mark_annual_leave_source_only(event, [msg])
    file = tmp_path / 'phep_nam.xlsx'
    export_to_excel([event], {'month':10,'year':2026,'snapshot':[msg]}, file)
    book = openpyxl.load_workbook(file)
    try:
        assert book['1_TIN_NHAN_GOC']['E5'].value == body
        assert book['1_TIN_NHAN_GOC']['F5'].value == 'Thông tin'
        assert book['1_TIN_NHAN_GOC']['H5'].value == 0
        assert book['2_CHAM_CONG']['A5'].value is None
        assert book['4_CAN_KIEM_TRA']['A5'].value is None
    finally:
        book.close()
