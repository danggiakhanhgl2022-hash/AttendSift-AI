"""Source-only policy must not erase actual attendance events or trust AI reasons."""
from copy import deepcopy

import pytest

from src.services.event_policy import is_information_only, mark_information_only


def message(body, message_id='m1', name='Người Giả Lập', day='2026-10-13'):
    return {'message_id': message_id, 'account_id': 101, 'name': name,
            'body': body, 'send_time': f'{day}T08:51:57+07:00'}


def attendance(**changes):
    event = {'employee_name': 'Người Giả Lập', 'person_id': 'model-person',
             'event_date': '2026-10-12', 'event_type': 'Đi trễ',
             'source_message_ids': ['m1'], 'needs_review': True,
             'review_state': 'pending', 'is_excluded': False,
             'review_reason': 'Chưa rõ chính sách tính công cho lịch bên ngoài.',
             'validation_errors': ['start_time: giờ không hợp lệ'],
             'validation_warnings': ['Cần kiểm tra chính sách'],
             'minutes_not_compensated': 150, 'minutes_compensated': 0,
             'ot_minutes': 0, 'ot_used_minutes': 0, 'leave_days': 0}
    event.update(changes)
    return event


def test_school_timetable_is_kept_as_source_without_attendance_or_review():
    snapshot = [message('Dạ em gửi lịch qua trường Kinh tế tuần này ạ\n'
                        '- Chiều 12/10 từ 14h30 đến 17h (chấm công ra lúc 14h30)\n'
                        '- Chiều 16/10 từ 14h30 đến 17h (chấm công ra lúc 14h30)')]
    source_before = deepcopy(snapshot)
    row = attendance()

    assert is_information_only(row, snapshot)
    assert mark_information_only(row, snapshot)
    assert row['source_only'] is True
    assert row['is_excluded'] is True
    assert row['review_state'] == 'excluded'
    assert row['needs_review'] is False
    assert row['minutes_not_compensated'] == 0
    assert row['minutes_compensated'] == 0
    assert row['ot_minutes'] == 0
    assert row['ot_used_minutes'] == 0
    assert row['leave_days'] is None
    assert snapshot == source_before
    assert row['event_date'] == '2026-10-12'
    assert row['employee_name'] == 'Người Giả Lập'


@pytest.mark.parametrize('body', [
    '3h15 em đi qua nhà hàng chuẩn bị với chị Giả Lập ạ',
    '3h15 chị/em đi qua nhà hàng chuẩn bị, có vấn đề cần xử lý gấp, mọi người điện cho chị/em.',
    'Em gửi lịch qua trường: ngày 1/10 8h-11h (check in lúc 11h30)',
])
def test_company_outing_or_school_schedule_without_attendance_claim_is_source_only(body):
    row = attendance(event_date='2026-10-01' if '1/10' in body else '2026-10-13')
    assert is_information_only(row, [message(body)])


@pytest.mark.parametrize('claim', [
    'Em xin phép về sớm 30 phút, sẽ làm bù vào ngày mai.',
    'Em xin phép nghỉ nửa ngày phép năm.',
    'Em xin phép đi trễ 30 phút không làm bù.',
    'Em tăng ca 30 phút hỗ trợ dự án.',
    'Giả Lập OT 4h hỗ trợ dự án.',
])
def test_outing_does_not_hide_explicit_leave_late_overtime_or_compensation(claim):
    snapshot = [message('Em đi qua nhà hàng chuẩn bị.\n' + claim)]
    row = attendance(event_date='2026-10-13')
    before = deepcopy(row)
    assert not is_information_only(row, snapshot)
    assert not mark_information_only(row, snapshot)
    assert row == before


@pytest.mark.parametrize('minutes', [30, 240])
def test_valid_overtime_remains_attendance(minutes):
    snapshot = [message(f'Giả Lập OT {minutes} phút ngày 13/10 hỗ trợ dự án.')]
    row = attendance(event_date='2026-10-13', event_type='Tăng ca',
                     minutes_not_compensated=0, ot_minutes=minutes,
                     needs_review=False, validation_errors=[], validation_warnings=[])
    assert not is_information_only(row, snapshot)
    assert not mark_information_only(row, snapshot)
    assert row['ot_minutes'] == minutes
    assert not row['is_excluded']


def test_ai_reason_does_not_prove_business_outing():
    snapshot = [message('Em có việc cá nhân, xin nghỉ nửa ngày phép năm.')]
    row = attendance(event_date='2026-10-13', event_type='Nghỉ', leave_days=0.5,
                     reason='Đi công tác bên ngoài',
                     review_reason='Đi ra ngoài làm việc công ty, chưa rõ chính sách trừ công.')
    assert not is_information_only(row, snapshot)


@pytest.mark.parametrize('source_ids', [[], ['missing'], ['m1', 'missing']])
def test_invalid_source_links_cannot_automatically_exclude_an_event(source_ids):
    snapshot = [message('Em đi qua nhà hàng chuẩn bị.')]
    row = attendance(event_date='2026-10-13', source_message_ids=source_ids)
    assert not is_information_only(row, snapshot)


def test_unrelated_information_message_in_snapshot_is_not_used_as_evidence():
    snapshot = [message('Em xin nghỉ nửa ngày phép năm.', message_id='attendance'),
                message('Em qua nhà hàng chuẩn bị.', message_id='outing')]
    row = attendance(event_date='2026-10-13', source_message_ids=['attendance'])
    assert not is_information_only(row, snapshot)


def test_no_attendance_claim_is_source_only_even_if_ai_hallucinates_event_date():
    snapshot = [message('Lịch qua trường Kinh tế chiều 12/10 từ 14h30 đến 17h.')]
    row = attendance(event_date='2026-10-08')
    assert is_information_only(row, snapshot)


@pytest.mark.parametrize('day', [None, '', '2026-02-30'])
def test_non_event_cannot_become_payroll_row_even_with_bad_ai_date(day):
    snapshot = [message('Em qua trường Kinh tế chiều 12/10 từ 14h30 đến 17h.')]
    row = attendance(event_date=day)
    assert is_information_only(row, snapshot)


def test_unsupported_employee_cannot_create_an_event_without_timekeeping_claim():
    snapshot = [message('Em qua trường Kinh tế chiều 12/10 từ 14h30 đến 17h.')]
    row = attendance(employee_name='Một Người Khác')
    assert is_information_only(row, snapshot)


def test_information_only_export_preserves_source_without_attendance_or_review_rows(tmp_path):
    from openpyxl import load_workbook
    from src.utils.excel_export import export_to_excel

    snapshot = [message('Em qua nhà hàng chuẩn bị lúc 3h15 ạ.')]
    row = attendance(event_date='2026-10-13', event_id='info')
    assert mark_information_only(row, snapshot)
    path = tmp_path / 'source-only.xlsx'
    export_to_excel([row], {'month': 10, 'year': 2026, 'snapshot': snapshot}, path)
    workbook = load_workbook(path)
    try:
        assert workbook['1_TIN_NHAN_GOC']['A5'].value == 'm1'
        assert workbook['1_TIN_NHAN_GOC']['E5'].value == snapshot[0]['body']
        assert workbook['1_TIN_NHAN_GOC']['H5'].value == 0
        for sheet in ['2_CHAM_CONG', '3_TONG_KET', '4_CAN_KIEM_TRA']:
            assert workbook[sheet]['A5'].value is None
    finally:
        workbook.close()


def test_export_service_filters_false_attendance_and_keeps_real_overtime(temp_db, tmp_path):
    from pathlib import Path
    from unittest.mock import Mock, patch
    from openpyxl import load_workbook
    from src.models.database import (ConfigRepository, MessageRepository,
                                     ExportSessionRepository, ActionEventRepository)
    from src.services.export_service import ExportService

    config = ConfigRepository(temp_db)
    config.set('chatwork_room_id', 'synthetic-room')
    config.set('ai_api_key', 'fake-development-key')
    repository = MessageRepository(temp_db)
    messages = [
        message('Em gửi lịch qua trường Kinh tế: chiều 12/10 từ 14h30 đến 17h.',
                message_id='school'),
        message('3h15 em đi qua nhà hàng chuẩn bị với chị giả lập ạ.',
                message_id='outing', day='2026-10-09'),
        message('Người Giả Lập OT 4h ngày 8/10, hỗ trợ dự án giả lập.',
                message_id='ot-four-hours', day='2026-10-09'),
        message('Người Giả Lập OT 30ph ngày 13/10, hỗ trợ dự án giả lập.',
                message_id='ot-half-hour', day='2026-10-14'),
    ]
    for row in messages:
        row['room_id'] = 'synthetic-room'
    repository.save_messages_batch(messages)
    source_before = deepcopy(repository.get_by_room('synthetic-room'))
    service = ExportService(temp_db, config, repository,
                            ExportSessionRepository(temp_db), ActionEventRepository(temp_db),
                            str(tmp_path))
    baseline = {'identity_evidence': 'Người Giả Lập', 'identity_uncertain': False,
                'confidence': 0.95, 'event_subtype': None, 'start_time': None,
                'end_time': None, 'ot_credit_date': None, 'leave_days': None,
                'reason': 'Dữ liệu giả lập', 'hours_compensation': 'Không áp dụng'}
    untrusted = [
        attendance(**baseline, source_message_ids=['school']),
        attendance(**baseline, source_message_ids=['outing'], event_date='2026-10-09',
                   event_type='Làm bù', minutes_compensated=150),
        attendance(**baseline, source_message_ids=['ot-four-hours'], event_date='2026-10-08',
                   event_type='Tăng ca', minutes_not_compensated=0, ot_minutes=240,
                   needs_review=False, review_reason=None),
        attendance(**baseline, source_message_ids=['ot-half-hour'], event_date='2026-10-13',
                   event_type='Tăng ca', minutes_not_compensated=0, ot_minutes=30,
                   needs_review=False, review_reason=None),
    ]
    original_ai = deepcopy(untrusted)
    adapter = Mock()
    adapter.analyze_attendance.return_value = deepcopy(untrusted)
    session_id = service.create_session(10, 2026, accept_risk=True)['session_id']
    with patch('src.services.export_service.create_ai_adapter', return_value=adapter):
        assert service.analyze(session_id)['status'] == 'ready'

    assert service.get_results_for_review(session_id) == []
    kept = service.get_results(session_id)
    assert len(kept) == 2
    assert sorted(row['ot_minutes'] for row in kept) == [30, 240]
    assert all(row['review_state'] == 'valid' for row in kept)
    assert all(not row['validation_errors'] for row in kept)
    audited = service.get_results(session_id, include_excluded=True)
    information = [row for row in audited if row.get('source_only')]
    assert len(information) == 2
    assert all(row['review_state'] == 'excluded' and not row['needs_review']
               for row in information)
    assert all(row['minutes_compensated'] == row['minutes_not_compensated'] == 0
               for row in information)
    originals_by_source = {tuple(row['source_message_ids']): row for row in original_ai}
    for row in audited:
        assert row['ai_raw_output'] == originals_by_source[tuple(row['source_message_ids'])]
    assert repository.get_by_room('synthetic-room') == source_before
    assert service.verify_source_integrity(session_id)['valid']

    exported = service.export_to_excel(session_id)
    assert exported['event_count'] == 2
    assert Path(exported['file_path']).is_file()
    workbook = load_workbook(exported['file_path'])
    try:
        originals = workbook['1_TIN_NHAN_GOC']
        bodies = {originals.cell(index, 1).value: originals.cell(index, 5).value
                  for index in range(5, 9)}
        assert bodies == {row['message_id']: row['body'] for row in messages}
        detail = workbook['2_CHAM_CONG']
        assert [detail['T5'].value, detail['T6'].value] == [240, 30]
        assert sum(detail.cell(index, 22).value for index in [5, 6]) == 270
        assert detail['A7'].value is None
        summary = workbook['3_TONG_KET']
        assert summary['A5'].value == 'Người Giả Lập'
        assert summary['A6'].value is None
        assert workbook['4_CAN_KIEM_TRA']['A5'].value is None
    finally:
        workbook.close()



@pytest.mark.parametrize('body', [
    '9h30 chị/em sẽ đi ra ngoài có việc của công ty, có việc gì gấp gọi điện nhé!',
    'Từ 3h30 chiều chị/em đi ra ngoài có việc công ty.',
    'Dạ em gửi lịch qua trường Kinh tế tuần này ạ. Chiều 12/10 từ 14h30 đến 17h (chấm công ra lúc 14h30)',
    '3h15 em đi qua nhà hàng chuẩn bị với chị Uyên ạ',
    'Hôm nay là ngày làm việc cuối cùng của em. Cảm ơn mọi người!',
    'Ngày mai không có nghỉ và không làm bù, mọi người làm theo lịch bình thường.',
    'Em sẽ làm bù nếu công ty sắp xếp, hiện chưa thực hiện bù giờ.',
])
def test_information_only_realistic_chatwork_messages(body):
    row = attendance(event_date='2026-10-13', event_type='Nghỉ')
    assert is_information_only(row, [message(body)])
    assert mark_information_only(row, [message(body)])


@pytest.mark.parametrize('body', [
    'Em xin phép nghỉ chiều ngày 8/10 (không lương), em đã xin phép rồi.',
    'Em xin phép nghỉ sáng ngày 30/9 (nửa phép năm), lý do lên trường.',
    'Em xin phép off sáng T6 16/10 (0.5 phép năm).',
    'Em xin phép về sớm 30 phút không làm bù.',
    'Em xin phép đi muộn 30 phút không làm bù.',
    'Em xin phép đi trễ 30 phút, sẽ làm bù vào ngày mai.',
    'Hằng OT 30ph ngày 10/5 hỗ trợ khách.',
    'Quân tăng ca 4h ngày 8/10.',
    'Em xin phép nghỉ bù OT 30 phút chiều nay.',
])
def test_actual_attendance_claims_are_not_silently_excluded(body):
    row = attendance(event_date='2026-10-13')
    assert not is_information_only(row, [message(body)])




def test_month_without_any_attendance_still_exports_source_messages(temp_db, tmp_path):
    """A month of mere announcements must not block Export Excel."""
    from pathlib import Path
    from unittest.mock import Mock, patch
    from openpyxl import load_workbook
    from src.models.database import (
        ConfigRepository, MessageRepository, ExportSessionRepository,
        ActionEventRepository,
    )
    from src.services.export_service import ExportService

    config = ConfigRepository(temp_db)
    config.set('chatwork_room_id', 'synthetic-room')
    config.set('ai_api_key', 'fake-test-key')
    repo = MessageRepository(temp_db)
    announcement = message(
        '3h15 em đi qua nhà hàng chuẩn bị cùng chị Uyên ạ.',
        message_id='only-message', day='2026-10-09'
    )
    announcement['room_id'] = 'synthetic-room'
    repo.save_messages_batch([announcement])
    service = ExportService(
        temp_db, config, repo, ExportSessionRepository(temp_db),
        ActionEventRepository(temp_db), str(tmp_path)
    )
    sid = service.create_session(10, 2026, accept_risk=True)['session_id']
    adapter = Mock()
    adapter.analyze_attendance.return_value = []
    with patch('src.services.export_service.create_ai_adapter', return_value=adapter):
        assert service.analyze(sid)['status'] == 'ready'
    assert service.get_results_for_review(sid) == []
    exported = service.export_to_excel(sid)
    assert exported['event_count'] == 0
    assert Path(exported['file_path']).is_file()
    workbook = load_workbook(exported['file_path'])
    try:
        assert workbook['1_TIN_NHAN_GOC']['A5'].value == 'only-message'
        assert workbook['1_TIN_NHAN_GOC']['E5'].value == announcement['body']
        assert workbook['2_CHAM_CONG']['A5'].value is None
        assert workbook['3_TONG_KET']['A5'].value is None
        assert workbook['4_CAN_KIEM_TRA']['A5'].value is None
    finally:
        workbook.close()
