"""Safe sample-only acceptance tests for the GitHub release."""
from io import BytesIO
import openpyxl
import pytest
from src.services.source_reconciliation import reconcile_explicit_records
from src.services.event_policy import is_ctv_text_only, configured_ctv_aliases
from src.utils.excel_export import export_to_excel


def _msg(message_id, body, sender='Người báo hộ'):
    return {'message_id': message_id, 'account_id': 1, 'name': sender,
            'body': body, 'send_time': '2026-10-08T09:00:00+07:00'}


def test_two_days_three_people_each_day_are_six_independent_ot_rows():
    body = ('Team A: ngày 5/10\nThư OT 1h, anh Tâm + Đăng OT 2h\n'
            'Ngày 7/10\nThư OT 1h30\nAnh Tâm OT 30ph\nViệt OT 1h')
    snapshot = [_msg('m1', body)]
    events = []
    reconcile_explicit_records(events, snapshot, 10, 2026)
    records = {(e['employee_name'], e['event_date'], e['ot_minutes'])
               for e in events if not e.get('is_excluded')}
    assert records == {('Thư','2026-10-05',60),('Tâm','2026-10-05',120),
                       ('Đăng','2026-10-05',120),('Thư','2026-10-07',90),
                       ('Tâm','2026-10-07',30),('Việt','2026-10-07',60)}
    output = BytesIO()
    export_to_excel(events, {'month': 10, 'year': 2026, 'snapshot': snapshot}, output)
    output.seek(0)
    wb = openpyxl.load_workbook(output)
    try:
        assert wb.sheetnames == ['1_TIN_NHAN_GOC','2_CHAM_CONG','3_TONG_KET','4_CAN_KIEM_TRA']
        assert wb['1_TIN_NHAN_GOC']['E5'].value == body
        rows = list(wb['2_CHAM_CONG'].iter_rows(min_row=5, max_row=10,values_only=True))
        assert len(rows) == 6
        assert all(row[25] == body for row in rows)
    finally: wb.close()


def test_bare_names_are_not_implicitly_ctv(monkeypatch):
    monkeypatch.delenv('CHAMCONG_CTV_ALIASES', raising=False)
    names = ['Cầm', 'Ny', 'Tráng']
    source = [_msg('m1', 'Team A: ngày 8/10\nCầm, Ny, Tráng OT 1h\nAn (CTV) OT 1h')]
    events = []
    reconcile_explicit_records(events, source, 10, 2026)
    kept = {e['employee_name'] for e in events if not e.get('is_excluded')}
    assert kept == set(names)
    assert is_ctv_text_only({'employee_name':'An (CTV)', 'source_message_ids':['m1']}, source)


def test_explicit_ctv_aliases_are_optional_and_invalid_config_rejected(monkeypatch):
    monkeypatch.setenv('CHAMCONG_CTV_ALIASES','["Ny"]')
    assert configured_ctv_aliases() == {'ny'}
    assert is_ctv_text_only({'employee_name':'Ny'}, [])
    monkeypatch.setenv('CHAMCONG_CTV_ALIASES','[1]')
    with pytest.raises(ValueError, match='CHAMCONG_CTV_ALIASES'):
        configured_ctv_aliases()


def test_annual_leave_stays_source_only():
    body = 'Em xin phép nghỉ sáng 16/10 (0.5 phép năm)'
    messages = [_msg('m1',body)]
    events = []
    reconcile_explicit_records(events, messages, 10, 2026)
    assert not [e for e in events if not e.get('is_excluded')]
