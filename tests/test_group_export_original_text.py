"""Verbatim Chatwork export and An (CTV) text-only policy regression tests."""
from copy import deepcopy
from pathlib import Path

import openpyxl

from src.services.source_reconciliation import reconcile_explicit_records
from src.utils.excel_export import export_to_excel
from tests.test_source_reconciliation import msg, row

TEMPLATE = Path(__file__).resolve().parents[1] / 'templates' / 'template.xlsx'


def _render(tmp_path, events, source):
    output = tmp_path / 'attendance.xlsx'
    for index, event in enumerate(events, 1):
        event['event_id'] = 'EVT%05d' % index
    export_to_excel(events, {'month': 10, 'year': 2026, 'snapshot': source},
                    output, template_path=TEMPLATE)
    return openpyxl.load_workbook(output)


def test_group_ot_rows_share_exact_unchanged_chatwork_content(tmp_path):
    original = ('Team DemoAlpha: ngày 01/10\n'
                'Tân, Hào, Quỳnh Anh OT 30 phút\n'
                'An (CTV) OT 30 phút\n'
                'Hỗ trợ dự án DemoAlpha. \n'
                '[info]Nguyên văn: «đối ứng»  日本語  & <tag>[/info]')
    source = [msg('m1', 'Lê Minh Tân', original, '2026-10-02', 101)]
    events = [row('m1', 'An (CTV)', '2026-10-01', 'Tăng ca', 30)]
    unreconciled = deepcopy(events)
    reconcile_explicit_records(events, source, 10, 2026)
    assert events[0]['is_excluded'] and events[0]['source_only']
    assert unreconciled[0]['ot_minutes'] == events[0]['ot_minutes'] == 30  # audit kept
    active = [e for e in events if not e.get('is_excluded')]
    assert {e['employee_name'] for e in active} == {'Lê Minh Tân', 'Hào', 'Quỳnh Anh'}
    assert all(e['ot_minutes'] == 30 for e in active)
    assert source[0]['body'] == original  # parser must never change the snapshot
    workbook = _render(tmp_path, events, source)
    try:
        assert workbook['1_TIN_NHAN_GOC']['E5'].value == original
        assert workbook['1_TIN_NHAN_GOC']['H5'].value == 3
        sheet = workbook['2_CHAM_CONG']
        assert {sheet.cell(i, 6).value for i in range(5, 8)} == {e['employee_name'] for e in active}
        assert [sheet.cell(i, 20).value for i in range(5, 8)] == [30, 30, 30]
        assert [sheet.cell(i, 26).value for i in range(5, 8)] == [original] * 3
        assert sheet['A8'].value is None  # No An (CTV) payroll row
    finally:
        workbook.close()


def test_only_my_ctv_message_is_exported_as_original_text(tmp_path):
    body = 'An (CTV) OT 30 phút\nNguyên văn, không thay đổi.  '
    source = [msg('m2', 'Người báo hộ', body, '2026-10-02', 102)]
    events = [row('m2', 'An', '2026-10-01', 'Tăng ca', 30)]
    # A missing report date must not cause a An (CTV) AI row to leak into OT.
    original = source[0]['body']
    reconcile_explicit_records(events, source, 10, 2026)
    assert len(events) == 1 and events[0]['is_excluded']
    workbook = _render(tmp_path, events, source)
    try:
        assert workbook['1_TIN_NHAN_GOC']['E5'].value == original
        assert workbook['1_TIN_NHAN_GOC']['F5'].value == 'Thông tin'
        assert workbook['2_CHAM_CONG']['A5'].value is None
        assert workbook['3_TONG_KET']['A5'].value is None
    finally:
        workbook.close()
