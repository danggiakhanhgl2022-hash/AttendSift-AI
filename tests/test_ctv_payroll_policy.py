import openpyxl
from src.services.event_policy import is_ctv_text_only
from src.services.source_reconciliation import reconcile_explicit_records
from src.utils.excel_export import export_to_excel


def test_mixed_report_excludes_only_marked_ctv_and_keeps_unknown_names(tmp_path):
    source = [{'message_id': 'm1', 'account_id': 1, 'name': 'Reporter',
               'send_time': '2026-10-08T09:00:00+07:00',
               'body': 'Team DemoBeta: ngày 8/10\nBảo, Cầm (CTV), Ny OT 1h\nTráng OT 30ph'}]
    events = []
    reconcile_explicit_records(events, source, 10, 2026)
    active = [e for e in events if not e.get('is_excluded')]
    assert {e['employee_name'] for e in active} == {'Bảo', 'Ny', 'Tráng'}
    path = tmp_path / 'payroll.xlsx'
    export_to_excel(events, {'month': 10, 'year': 2026, 'snapshot': source}, path)
    workbook = openpyxl.load_workbook(path)
    assert '(CTV)' in workbook.worksheets[0]['E5'].value
    assert {row[5] for row in workbook.worksheets[1].iter_rows(min_row=5, values_only=True)} == {'Bảo', 'Ny', 'Tráng'}
    workbook.close()


def test_unmarked_cam_is_not_ctv_without_explicit_confirmation():
    event = {'employee_name': 'Cầm', 'source_message_ids': ['m1']}
    source = [{'message_id': 'm1', 'body': 'Cầm OT 1h\nBảo OT 1h\nAn (CTV) OT 1h'}]
    assert not is_ctv_text_only(event, source)
    assert not is_ctv_text_only({'employee_name': 'Bảo', 'source_message_ids': ['m1']}, source)


def test_unmarked_cam_report_keeps_both_attendance_rows_and_original_text(tmp_path):
    source = [{'message_id': 'm1', 'account_id': 1, 'name': 'Thư',
               'send_time': '2026-10-02T09:08:00+07:00',
               'body': 'Team DemoBeta: ngày 1/10\nThư OT 1h: đối ứng khách ClientSix\nCầm OT 1h: đối ứng ClientSix'}]
    events = []
    reconcile_explicit_records(events, source, 10, 2026)
    assert {e['employee_name'] for e in events if not e.get('is_excluded')} == {'Thư', 'Cầm'}
    output = tmp_path / 'cam.xlsx'
    export_to_excel(events, {'month': 10, 'year': 2026, 'snapshot': source}, output)
    workbook = openpyxl.load_workbook(output)
    assert workbook.worksheets[0]['E5'].value == source[0]['body']
    assert workbook.worksheets[1]['F5'].value == 'Thư'
    assert workbook.worksheets[1]['F6'].value == 'Cầm'
    assert workbook.worksheets[1]['T5'].value == 60
    assert workbook.worksheets[2]['A5'].value == 'Cầm' or workbook.worksheets[2]['A5'].value == 'Thư'
    assert workbook.worksheets[3]['D5'].value is None
    workbook.close()


def test_export_filters_old_saved_ctv_row_without_losing_coworker(tmp_path):
    source = [{'message_id': 'm1', 'name': 'Reporter', 'send_time': '2026-10-08T09:00:00+07:00',
               'body': 'Cầm (CTV) OT 1h\nBảo OT 1h'}]
    events = [{'event_id': name, 'person_id': name, 'employee_name': name,
               'event_date': '2026-10-08', 'event_type': 'Tăng ca', 'ot_minutes': 60,
               'source_message_ids': ['m1']} for name in ['Cầm', 'Bảo']]
    path = tmp_path / 'old-session.xlsx'
    export_to_excel(events, {'month': 10, 'year': 2026, 'snapshot': source}, path)
    workbook = openpyxl.load_workbook(path)
    assert workbook.worksheets[1]['F5'].value == 'Bảo'
    assert workbook.worksheets[1].max_row == 5
    assert workbook.worksheets[2]['A5'].value == 'Bảo'
    workbook.close()
