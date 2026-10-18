from io import BytesIO
import hashlib
import re
import openpyxl
import pytest
from src.app.server import create_app
from src.app.workbook_view import read_workbook_summary


def workbook_bytes(formula=False):
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name in ['1_TIN_NHAN_GOC','2_CHAM_CONG','3_TONG_KET','4_CAN_KIEM_TRA','5_CTV']:wb.create_sheet(name)
    detail=wb['2_CHAM_CONG'];detail.cell(5,1,'EVT01');detail.cell(5,6,'Nhân viên Demo')
    summary=wb['3_TONG_KET'];summary.cell(5,1,'Nhân viên Demo');summary.cell(5,9,'=60+60' if formula else 120);summary.cell(5,14,.5)
    ctv=wb['5_CTV'];ctv.cell(5,1,'EVT02');ctv.cell(5,2,'CTV Demo');ctv.cell(8,1,'TỔNG THEO NGƯỜI');ctv.cell(9,1,'CTV-1');ctv.cell(9,2,'CTV Demo');ctv.cell(9,3,30);ctv.cell(9,4,0)
    wb['4_CAN_KIEM_TRA'].cell(5,6,'Cần kiểm tra (chưa tính công)')
    stream=BytesIO();wb.save(stream);wb.close();return stream.getvalue()


def test_groups_come_from_excel_sheets_and_exact_stored_values():
    value=read_workbook_summary(workbook_bytes(),'chamcong.xlsx')
    assert [(p['name'],p['classification'],p['ot'],p['leave']) for p in value['people']]==[('Nhân viên Demo','official',120,.5),('CTV Demo','contractor',30,0)]
    assert value['pending']==1


def test_uncached_formula_is_unknown_not_zero():
    value=read_workbook_summary(workbook_bytes(True),'chamcong.xlsx')
    assert value['people'][0]['ot'] is None and value['missing_cached_values']


@pytest.mark.parametrize('content,name',[(b'invalid','x.xlsx'),(b'invalid','x.csv')])
def test_reject_invalid_workbook(content,name):
    with pytest.raises(ValueError):read_workbook_summary(content,name)


def test_upload_is_read_only_and_protected(tmp_path):
    app=create_app(tmp_path);client=app.test_client();services=app.extensions['services']
    before=services['sessions'].get_session('NONE')
    payload=workbook_bytes();token=re.search(r'name="csrf-token" content="([^"]+)"',client.get('/').get_data(as_text=True))[1]
    assert client.post('/api/dashboard/workbook-summary',data={'file':(BytesIO(payload),'demo.xlsx')}).status_code==403
    result=client.post('/api/dashboard/workbook-summary',headers={'X-CSRF-Token':token},data={'file':(BytesIO(payload),'demo.xlsx')})
    assert result.status_code==200 and result.get_json()['data']['people'][0]['classification']=='official'
    assert services['sessions'].get_session('NONE')==before
    with services['db']._get_connection() as conn:
        assert conn.execute('select count(*) from ai_results').fetchone()[0]==0
        assert conn.execute('select count(*) from export_sessions').fetchone()[0]==0


def test_session_workbook_is_read_without_reclassifying_records(tmp_path):
    app=create_app(tmp_path);services=app.extensions['services'];sid='DEMO-SHEET'
    services['sessions'].create_session(sid,10,2026)
    path=tmp_path/'outputs'/'sessions'/sid/'demo.xlsx';path.parent.mkdir(parents=True);content=workbook_bytes();path.write_bytes(content)
    services['sessions'].update_session(sid,status='exported',file_path=str(path))
    result=app.test_client().get(f'/api/export/{sid}/workbook-summary')
    assert result.status_code==200 and result.get_json()['data']['people'][0]['ot']==120
    assert hashlib.sha256(path.read_bytes()).digest()==hashlib.sha256(content).digest()


def test_missing_session_workbook_has_no_invented_totals(tmp_path):
    result=create_app(tmp_path).test_client().get('/api/export/NONE/workbook-summary')
    assert result.status_code==200 and result.get_json()['data']=={'available':False}
