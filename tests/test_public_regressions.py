"""Public, synthetic checks of name-lot resolution and five-sheet contract."""
from io import BytesIO
from zipfile import ZipFile
from src.services.personnel_policy import candidates, classify
from src.services.source_reconciliation import extract_explicit_claims, reconcile_explicit_records
from src.services.validation import validate_set
from src.utils.excel_export import export_to_excel


def person(full, code):
    return {'id':code,'employee_code':code,'full_name':full,'aliases':[],
            'chatwork_account_id':'','active':True}


def roster():
    return {'initialized':True,'revision':1,'employees':[
        person('Nguyễn Minh Quân','NV001'),person('Trần Đức Quân','NV002'),
        person('Lê Thu Bình','NV003')]}


def message(body):
    return {'message_id':'demo1','room_id':'room','name':'Người báo cáo (Demo)',
            'account_id':'999','body':body,'send_time':'2026-10-16T08:00:00+07:00'}


def test_middle_name_disambiguates_same_given_name():
    people=roster()['employees']
    assert [p['id'] for p in candidates('Minh Quân',people)]==['NV001']
    assert [p['id'] for p in candidates('Đức Quân',people)]==['NV002']
    assert len(candidates('Quân',people))==2
    assert [p['id'] for p in candidates('Bình',people)]==['NV003']
    assert candidates('Người Ngoài',people)==[]


def test_split_first_then_classify_and_keep_reviewable_ambiguity():
    m=message('Team Demo: ngày 15/10\nMinh Quân + Đức Quân OT 2h\nQuân OT 1h\nLan OT 30p')
    claims=extract_explicit_claims(m,[m],10,2026,True)
    assert {(c['employee_name'],c['ot_minutes']) for c in claims} == {
        ('Minh Quân',120),('Đức Quân',120),('Quân',60),('Lan',30)}
    records=[]
    reconcile_explicit_records(records,[m],10,2026,roster_policy=True)
    assert len(records)==4
    groups={}
    for r in records:
        groups[r['source_employee_name']]=classify(r,roster(),[m])
    assert groups == {'Minh Quân':'official','Đức Quân':'official',
                      'Quân':'unresolved','Lan':'contractor'}


def test_five_sheet_export_has_source_official_review_and_contractor():
    m=message('Team Demo: ngày 15/10\nMinh Quân OT 2h\nĐức Quân OT 1h\nLan OT 30p')
    records=[]
    reconcile_explicit_records(records,[m],10,2026,roster_policy=True)
    for i,r in enumerate(records,1):
        r['event_id']=f'EVT{i:05d}'
        classify(r,roster(),[m])
    validate_set(records,[m],10,2026)
    assert len(records)==3
    output=BytesIO()
    export_to_excel(records,{'month':10,'year':2026,'snapshot':[m],
                             'ai_settings':{'employee_registry':roster()}},output)
    with ZipFile(BytesIO(output.getvalue())) as excel:
        workbook=excel.read('xl/workbook.xml').decode()
        assert all(f'name="{sheet}"' in workbook for sheet in
          ('1_TIN_NHAN_GOC','2_CHAM_CONG','3_TONG_KET','4_CAN_KIEM_TRA','5_CTV'))


def test_ambiguous_given_name_does_not_enter_official_totals():
    import openpyxl
    m=message('Team Demo: ngày 15/10\nMinh Quân OT 2h\nĐức Quân OT 1h\nQuân OT 30p\nLan OT 1h')
    records=[]
    reconcile_explicit_records(records,[m],10,2026,roster_policy=True)
    assert len(records)==4
    for i,r in enumerate(records,1):
        r['event_id']=f'EVT{i:05d}'
        classify(r,roster(),[m])
    validate_set(records,[m],10,2026)
    out=BytesIO()
    export_to_excel(records,{'month':10,'year':2026,'snapshot':[m],
            'ai_settings':{'employee_registry':roster()}},out)
    wb=openpyxl.load_workbook(BytesIO(out.getvalue()),read_only=True,data_only=True)
    official=[str(v) for r in wb['2_CHAM_CONG'].iter_rows(values_only=True) for v in r if v is not None]
    review=[str(v) for r in wb['4_CAN_KIEM_TRA'].iter_rows(values_only=True) for v in r if v is not None]
    ctv=[str(v) for r in wb['5_CTV'].iter_rows(values_only=True) for v in r if v is not None]
    assert 'Nguyễn Minh Quân' in official and 'Trần Đức Quân' in official
    assert 'Quân' in review
    assert 'Lan' in ctv
