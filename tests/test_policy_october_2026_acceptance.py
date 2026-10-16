"""Acceptance cases copied from October 2026 user reports.

Contract: deterministic group OT becomes a validated row for every literal
person/day/duration, without a registry requirement; annual leave/An (CTV) are
raw Chatwork text only. Model hallucinations never enter payroll.
"""
from copy import deepcopy
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl
import pytest

from src.models.database import (Database, ConfigRepository, MessageRepository,
                                 ExportSessionRepository, ActionEventRepository)
from src.services.export_service import ExportService
from src.services.identity_context import repair_identity_evidence, align_event_identities
from src.services.source_reconciliation import reconcile_explicit_records
from src.services.validation import validate_set
from src.utils.excel_export import export_to_excel
from tests.test_source_reconciliation import msg, row

TEMPLATE = Path(__file__).resolve().parents[1] / 'templates' / 'template.xlsx'


def _process(messages, model_rows=None):
    rows = deepcopy(model_rows or [])
    reconcile_explicit_records(rows, messages, 10, 2026)
    for event in rows:
        if not event.get('is_excluded'):
            repair_identity_evidence(event, messages)
    align_event_identities([r for r in rows if not r.get('is_excluded')], messages)
    validate_set([r for r in rows if not r.get('is_excluded')], messages, 10, 2026)
    return rows


def _active(rows):
    return [r for r in rows if not r.get('is_excluded')]


@pytest.mark.parametrize('body,expected', [
    ('Team DemoBeta: ngày 5/10\nThư OT 1h, anh Tâm + Đăng OT 2h\nđối ứng khách ClientFour\n\n'
     'Ngày 7/10\nThư OT 1h30: đối ứng ClientFour\nAnh Tâm OT 30ph\nViệt OT 1h',
     {('Trần An Thư',5,60),('Tâm',5,120),('Đăng',5,120),
      ('Trần An Thư',7,90),('Tâm',7,30),('Việt',7,60)}),
    ('Team DemoAlpha: ngày 05/10\nTân, Hằng, Quỳnh Anh OT 30ph\nHỗ trợ dự án DemoAlpha',
     {('Lê Minh Tân',5,30),('Hằng',5,30),('Quỳnh Anh',5,30)}),
    ('Team DemoAlpha: ngày 09/10\nTân, Quỳnh Anh, Hằng, Duy OT 30ph',
     {('Lê Minh Tân',9,30),('Quỳnh Anh',9,30),('Hằng',9,30),('Duy',9,30)}),
    ('Team DemoBeta: ngày 8/10\nThư OT 4h: đối ứng khách ClientSix\n'
     'anh Tâm OT 30ph: đối ứng ClientOne\nĐăng, Bảo, Quân, Việt OT 4h',
     {('Trần An Thư',8,240),('Tâm',8,30),('Đăng',8,240),
      ('Bảo',8,240),('Quân',8,240),('Việt',8,240)}),
    ('Team DemoBeta: ngày 13/10\nThư OT 2h: đối ứng khách SAC\n'
     'Ny OT 2h: đối ứng SAC\nCẩm OT 2h: đối ứng SOA\n'
     'anh Tâm, Việt, Bảo OT 1h30: đến cửa hàng DemoBeta',
         {('Trần An Thư',13,120),('Ny',13,120),('Cẩm',13,120),
      ('Tâm',13,90),('Việt',13,90),('Bảo',13,90)}),
    ('Team DemoBeta: ngày 9/10\nanh Tâm OT 30ph\nNgày 10/10:\nBảo OT 8h\nViệt OT 4h\n'
     'Ngày 12/10:\nThư OT 1h\nanh Tâm OT 1h\nViệt OT 1h\nTrắng OT 30ph',
     {('Tâm',9,30),('Bảo',10,480),('Việt',10,240),('Trần An Thư',12,60),
      ('Tâm',12,60),('Việt',12,60),('Trắng',12,30)}),
])
def test_actual_reports_split_into_valid_payroll_without_registry_warnings(body, expected):
    author = 'Lê Minh Tân' if 'DemoAlpha' in body else 'Trần An Thư'
    source = [msg('report',author,body,'2026-10-14',101)]
    rows = _active(_process(source))
    assert {(r['employee_name'],int(r['event_date'][-2:]),r['ot_minutes']) for r in rows} == expected
    assert all(not r['needs_review'] and not r['validation_errors'] and
               not r['validation_warnings'] for r in rows)
    assert all(r['source_message_ids'] == ['report'] for r in rows)
    assert source[0]['body'] == body


def test_bad_model_alias_missing_id_and_wrong_date_do_not_block_group_overtime():
    source = [msg('m','Trần An Thư', 'Team DemoBeta: ngày 8/10\nThư OT 4h\n'
                  'anh Tâm OT 30ph\nĐăng, Bảo, Quân, Việt OT 4h','2026-10-09',13)]
    ai = [row('m','Đăng','2026-10-08','Tăng ca',240,
              confidence=.6, identity_uncertain=True,needs_review=True,
              review_reason='Chưa rõ mã nhân viên cho Đăng.'),
          row('m','Đăng','2026-10-09','Tăng ca',240,needs_review=True,
              review_reason='Ngày AI chưa có bằng chứng.'),
          row('m','Báo','2026-10-08','Tăng ca',240)]
    rows=_process(source,ai)
    assert len(_active(rows)) == 6
    assert all(not r['needs_review'] and not r['validation_errors'] and not r['validation_warnings']
               for r in _active(rows))
    assert {(r['employee_name'],r['event_date']) for r in rows if r.get('is_excluded')} == {
        ('Đăng','2026-10-09'),('Báo','2026-10-08')}
    assert next(r for r in _active(rows) if r['employee_name']=='Đăng')['ot_minutes']==240


def test_annual_leave_without_model_row_is_in_original_sheet_only(tmp_path):
    body = ('Em xin phép nghỉ ngày 7-9/10 (thứ tư - thứ sáu) và ngày 12/10 (phép năm)\n'
            'Lý do: Em có việc cá nhân.\nEm xin cảm ơn ạ! ')
    src = [msg('annual','Võ Ngọc Quỳnh',body,'2026-09-25',66)]
    rows = _process(src)
    assert rows == []
    output = tmp_path/'leave.xlsx'
    export_to_excel(rows, {'month':10,'year':2026,'snapshot':src},output,TEMPLATE)
    wb=openpyxl.load_workbook(output)
    try:
        assert wb['1_TIN_NHAN_GOC']['E5'].value == body
        assert wb['1_TIN_NHAN_GOC']['F5'].value == 'Thông tin'
        assert wb['2_CHAM_CONG']['A5'].value is None
        assert wb['3_TONG_KET']['A5'].value is None
    finally:
        wb.close()


def test_bad_ai_annual_leave_not_counted_and_original_is_preserved(tmp_path):
    body='Em xin phép nghỉ sáng 5/10 (1/2 ngày phép năm)\nEm xin cảm ơn!'
    src=[msg('m','Bùi Ngân Hoa',body,'2026-10-01',22)]
    raw=row('m','Bùi Ngân Hoa','2026-10-05','Nghỉ',leave_days=.5)
    rows=_process(src,[raw])
    assert len(rows)==1 and rows[0]['source_only'] and rows[0]['is_excluded']
    output=tmp_path/'annual.xlsx'
    export_to_excel(rows,{'month':10,'year':2026,'snapshot':src},output,TEMPLATE)
    wb=openpyxl.load_workbook(output)
    try:
        assert wb['1_TIN_NHAN_GOC']['E5'].value==body
        assert wb['1_TIN_NHAN_GOC']['F5'].value=='Thông tin'
        assert wb['2_CHAM_CONG']['A5'].value is None
    finally:
        wb.close()


def test_report_with_my_ctv_and_three_ot_staff_keeps_raw_text(tmp_path):
    body='Team DemoAlpha: ngày 01/10\nTân, Hào, Quỳnh Anh OT 30 phút\nAn (CTV) OT 30 phút'
    source=[msg('p','Lê Minh Tân',body,'2026-10-02',19)]
    rows=_process(source,[row('p','An (CTV)','2026-10-01','Tăng ca',30)])
    assert [(r['employee_name'],r['ot_minutes']) for r in _active(rows)] == [
        ('Lê Minh Tân',30),('Hào',30),('Quỳnh Anh',30)]
    output=tmp_path/'ot.xlsx'
    export_to_excel(rows,{'month':10,'year':2026,'snapshot':source},output,TEMPLATE)
    wb=openpyxl.load_workbook(output)
    try:
        assert wb['1_TIN_NHAN_GOC']['E5'].value==body
        assert [wb['2_CHAM_CONG'].cell(i,26).value for i in (5,6,7)] == [body]*3
        assert wb['2_CHAM_CONG']['A8'].value is None
    finally:wb.close()


def test_stable_short_name_id_over_two_team_reports():
    sources=[msg('a','Lê Minh Tân', 'Team DemoAlpha: ngày 05/10\nTân, Hằng OT 30ph','2026-10-05',1),
             msg('b','Lê Minh Tân', 'Team DemoAlpha: ngày 09/10\nTân, Hằng OT 30ph','2026-10-09',1)]
    rows=_active(_process(sources))
    hang=[r for r in rows if r['employee_name']=='Hằng']
    assert len(hang)==2
    assert hang[0]['person_id']==hang[1]['person_id']
    assert all(not r['validation_warnings'] for r in rows)


def test_short_alias_with_multiple_sender_namesakes_is_not_a_conflict_in_report():
    report=msg('r','Người tổng hợp','Team DemoBeta: ngày 08/10\nTâm OT 30ph','2026-10-09',10)
    sources=[report,msg('a','Người giả lập Tâm','', '2026-10-01',11),
            msg('b','Trần Tâm','', '2026-10-02',12)]
    rows=_active(_process(sources))
    assert len(rows)==1 and rows[0]['employee_name']=='Tâm'
    assert rows[0]['person_id'].startswith('SRC-N-')
    assert not rows[0]['needs_review'] and not rows[0]['validation_warnings']


def test_unparseable_group_line_must_not_auto_discard_model_row():
    body='Team DemoBeta: ngày 08/10\nThư OT 1h\nĐăng OT thời lượng chưa rõ'
    source=[msg('u','Người báo hộ',body,'2026-10-09',33)]
    ai=row('u','Đăng','2026-10-08','Tăng ca',60,needs_review=True,review_reason='Thời lượng chưa rõ')
    rows=_process(source,[ai])
    assert not rows[0].get('is_excluded') and rows[0]['needs_review']
    assert any(r['employee_name']=='Thư' for r in rows)


def test_annual_leave_only_message_with_fabricated_overtime_is_excluded():
    body='Em xin phép nghỉ ngày 05/10 (phép năm)'
    src=[msg('leave','Người A',body,'2026-10-01',13)]
    fabricated=row('leave','Người A','2026-10-05','Tăng ca',120)
    rows=_process(src,[fabricated])
    assert len(rows)==1 and rows[0]['is_excluded']


def test_one_message_to_excel_individual_ot_same_raw_text_all_staff(tmp_path):
    body='Team DemoBeta: ngày 8/10\nThư OT 4h\nanh Tâm OT 30ph\nĐăng, Bảo, Quân, Việt OT 4h'
    src=[msg('p','Trần An Thư',body,'2026-10-09',14)]
    rows=_active(_process(src))
    for i,r in enumerate(rows,1):r['event_id']=f'EVT{i:05d}'
    path=tmp_path/'out.xlsx'
    export_to_excel(rows,{'month':10,'year':2026,'snapshot':src},path,TEMPLATE)
    wb=openpyxl.load_workbook(path)
    try:
        assert wb['1_TIN_NHAN_GOC']['E5'].value==body
        sheet=wb['2_CHAM_CONG']
        assert len(rows)==6
        assert [sheet.cell(i,20).value for i in range(5,11)].count(30)==1
        assert [sheet.cell(i,20).value for i in range(5,11)].count(240)==5
        assert [sheet.cell(i,26).value for i in range(5,11)] == [body]*6
    finally:wb.close()


def test_full_sqlite_analysis_and_excel_export_with_mixed_reports(tmp_path):
    db=Database(str(tmp_path/'data.db'))
    cfg=ConfigRepository(db)
    cfg.set('chatwork_room_id','room')
    repo=MessageRepository(db)
    bodies = {
        'group': 'Team DemoBeta: ngày 8/10\nThư OT 4h\nanh Tâm OT 30ph\nĐăng, Bảo, Quân, Việt OT 4h',
        'leave': 'Em xin phép nghỉ 7-9/10 và 12/10 (phép năm)\nEm xin cảm ơn ạ!',
        'my': 'Team DemoAlpha: ngày 01/10\nTân, Hào, Quỳnh Anh OT 30ph\nAn (CTV) OT 30ph',
    }
    repo.save_messages_batch([
        {'message_id':'group','room_id':'room','account_id':11,'name':'Trần An Thư',
         'body':bodies['group'],'send_time':'2026-10-09T09:00:00+07:00'},
        {'message_id':'leave','room_id':'room','account_id':12,'name':'Võ Ngọc Quỳnh',
         'body':bodies['leave'],'send_time':'2026-09-25T10:00:00+07:00'},
        {'message_id':'my','room_id':'room','account_id':13,'name':'Lê Minh Tân',
         'body':bodies['my'],'send_time':'2026-10-02T09:00:00+07:00'},
    ])
    service=ExportService(db,cfg,repo,ExportSessionRepository(db),ActionEventRepository(db),str(tmp_path))
    ai=Mock()
    # Completely incomplete and mistaken model: reconciliation must be robust.
    ai.analyze_attendance.return_value=[
        row('group','Đăng','2026-10-09','Tăng ca',240,
            confidence=.52,needs_review=True,review_reason='Chưa rõ mã nhân viên cho Đăng'),
        row('leave','Võ Ngọc Quỳnh','2026-10-07','Nghỉ'),
        row('my','An (CTV)','2026-10-01','Tăng ca',30),
    ]
    fake_cfg={'provider':'custom','model':'fake','base_url':'https://example.invalid/v1','api_key':'offline'}
    with patch.object(cfg,'get_ai_provider',return_value=fake_cfg), patch(
        'src.services.export_service.create_ai_adapter',return_value=ai):
        sid=service.create_session(10,2026,accept_risk=True)['session_id']
        answer=service.analyze(sid)
    assert answer['status']=='ready'
    assert service.get_results_for_review(sid)==[]
    active=service.get_results(sid)
    assert len(active)==9
    assert not any(r['employee_name'] in {'An (CTV)','Võ Ngọc Quỳnh'} for r in active)
    result=service.export_to_excel(sid)
    wb=openpyxl.load_workbook(result['file_path'])
    try:
        actual_raw={wb['1_TIN_NHAN_GOC'].cell(i,1).value:wb['1_TIN_NHAN_GOC'].cell(i,5).value
                    for i in range(5,8)}
        assert actual_raw == bodies
        attendance=wb['2_CHAM_CONG']
        assert {attendance.cell(i,20).value for i in range(5,14)}=={30,240}
        assert len([attendance.cell(i,6).value for i in range(5,14)])==9
        assert all(attendance.cell(i,26).value == bodies[attendance.cell(i,2).value]
                   for i in range(5,14))
        assert attendance['A14'].value is None
    finally:wb.close()


def test_one_chatwork_message_two_dates_six_independent_overtime_rows_and_original_excel(tmp_path):
    """User-confirmed DemoBeta example: OT is per employee, not divided across staff/days."""
    body = (
        'Team DemoBeta: ngày 5/10\n'
        'Thư OT 1h, anh Tâm + Đăng OT 2h\n'
        'đối ứng khách ClientFour\n\n'
        'Ngày 7/10\n'
        'Thư OT 1h30: Đối ứng khách ClientFour, ClientTwo, ClientFive, ClientSix, ClientEight\n'
        'Anh Tâm OT 30ph: Đối ứng khách ClientOne, ClientFour\n'
        'Việt OT 1h: đối ứng ClientEight'
    )
    source = [msg('two-days', 'Trần An Thư', body, '2026-10-08', 42)]
    results = _active(_process(source))
    expected = {
        ('Trần An Thư', '2026-10-05', 60),
        ('Tâm', '2026-10-05', 120),
        ('Đăng', '2026-10-05', 120),
        ('Trần An Thư', '2026-10-07', 90),
        ('Tâm', '2026-10-07', 30),
        ('Việt', '2026-10-07', 60),
    }
    assert len(results) == 6
    assert {(r['employee_name'], r['event_date'], r['ot_minutes']) for r in results} == expected
    assert all(not r['needs_review'] and not r['validation_warnings'] for r in results)
    assert len({r['person_id'] for r in results if r['employee_name']=='Tâm'}) == 1
    assert len({r['person_id'] for r in results if r['employee_name']=='Trần An Thư'}) == 1
    for idx, item in enumerate(results, 1):
        item['event_id'] = f'MULTI-{idx:03d}'
    output = tmp_path / 'multi_day.xlsx'
    export_to_excel(results, {'month': 10, 'year': 2026, 'snapshot': source}, output, TEMPLATE)
    wb = openpyxl.load_workbook(output)
    try:
        assert wb['1_TIN_NHAN_GOC']['E5'].value == body
        rows = wb['2_CHAM_CONG']
        actual = {(rows.cell(i, 6).value, rows.cell(i, 5).value.strftime('%Y-%m-%d'), rows.cell(i, 20).value)
                  for i in range(5, 11)}
        assert actual == expected
        assert all(rows.cell(i, 26).value == body for i in range(5, 11))
        assert rows.cell(11, 1).value is None
    finally:
        wb.close()


def test_08_oct_quan_literal_ot_is_valid_even_if_two_accounts_end_in_quan(tmp_path):
    """Screenshot regression: Quân 240 is not a duration conflict or HR-ID requirement."""
    body = ('Team DemoBeta: ngày 8/10\n'
            'Thư OT 4h: đối ứng khách ClientSix, ClientTwoの絵コンテ, ClientEight, ClientOne\n'
            'anh Tâm OT 30ph: đối ứng khách ClientOne\n'
            'Đăng, Bảo, Quân, Việt OT 4h: đối ứng ClientOne, ClientEight, ClientTwo, THREE')
    source = [msg('report','Trần An Thư',body,'2026-10-09',101),
              msg('q1','Phạm Văn Quân','', '2026-10-01',102),
              msg('q2','Bùi Minh Quân','', '2026-10-02',103)]
    ai = [row('report', 'Quân', '2026-10-08', 'Tăng ca',240,
              identity_uncertain=True,needs_review=True,
              review_reason='Cùng người/ngày có thời lượng mâu thuẫn hoặc có nhiều danh tính trùng tên.')]
    records = _active(_process(source, ai))
    assert {(r['employee_name'],r['ot_minutes']) for r in records} == {
        ('Trần An Thư', 240), ('Tâm', 30), ('Đăng', 240),
        ('Bảo', 240), ('Quân', 240), ('Việt', 240)}
    quan = next(r for r in records if r['employee_name'] == 'Quân')
    assert quan['person_id'].startswith('SRC-N-')  # not linked to either sender account
    assert not quan['identity_uncertain']
    assert not quan['needs_review']
    assert not quan['validation_errors'] and not quan['validation_warnings']
    assert all(not r['validation_errors'] and not r['validation_warnings'] for r in records)
    for i, r in enumerate(records, 1):
        r['event_id'] = f'EVT{i:05d}'
    path = tmp_path / 'quan.xlsx'
    export_to_excel(records,{'month':10,'year':2026,'snapshot':source},path,TEMPLATE)
    wb = openpyxl.load_workbook(path)
    try:
        assert wb['1_TIN_NHAN_GOC']['E5'].value == body
        events = wb['2_CHAM_CONG']
        assert len(records) == 6
        assert [events.cell(i,26).value for i in range(5,11)] == [body]*6
        assert [events.cell(i,20).value for i in range(5,11)].count(240) == 5
    finally:
        wb.close()


def test_actual_conflicting_minutes_for_same_person_and_day_requires_review():
    source = [msg('r','Người báo hộ','Team DemoBeta: ngày 08/10\nQuân OT 4h\nQuân OT 2h',
                  '2026-10-09',42)]
    rows = _active(_process(source))
    assert len(rows) == 1
    assert rows[0]['employee_name'] == 'Quân'
    assert rows[0]['needs_review']
    assert 'hai thời lượng OT khác nhau' in rows[0]['review_reason']
    assert rows[0]['validation_warnings']
