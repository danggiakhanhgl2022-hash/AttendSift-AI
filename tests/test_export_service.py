"""Behavioral acceptance tests using invented people and mocked AI only."""
import copy
import json
from pathlib import Path
import shutil
from unittest.mock import Mock,patch
import pytest
import openpyxl
from src.models.database import ConfigRepository,MessageRepository,ExportSessionRepository,ActionEventRepository
from src.services.export_service import ExportService,ExportServiceError
from src.services.validation import validate_set

TEMPLATE=Path(__file__).resolve().parents[1]/'templates'/'template.xlsx'


@pytest.fixture
def service(temp_db,tmp_path):
    config=ConfigRepository(temp_db)
    config.set('chatwork_room_id','test-room')
    config.set('ai_api_key','fake-development-key')
    messages=MessageRepository(temp_db)
    messages.save_messages_batch([{'message_id':'m1','room_id':'test-room','account_id':1,
        'name':'Người giả lập A','body':'Người giả lập A và Người giả lập B tăng ca 2 giờ ngày 05/09.',
        'send_time':'2026-09-05T18:00:00+07:00'}])
    (tmp_path/'templates').mkdir();shutil.copyfile(TEMPLATE,tmp_path/'templates'/'template.xlsx')
    return ExportService(temp_db,config,messages,ExportSessionRepository(temp_db),ActionEventRepository(temp_db),str(tmp_path))


@pytest.fixture
def event():
    return {'person_id':'P1','employee_name':'Người giả lập A','identity_evidence':'Người giả lập A',
        'identity_uncertain':False,'event_date':'2026-09-05','event_type':'Tăng ca','event_subtype':None,
        'reason':'Công việc giả lập','hours_compensation':'Không áp dụng','minutes_not_compensated':0,
        'minutes_compensated':0,'ot_minutes':120,'ot_used_minutes':0,'ot_credit_date':None,'leave_days':None,
        'start_time':None,'end_time':None,'source_message_ids':['m1'],'confidence':0.95,'needs_review':False,'review_reason':None}


def analyzed(service,event):
    sid=service.create_session(9,2026,accept_risk=True)['session_id']
    adapter=Mock();adapter.analyze_attendance.return_value=[copy.deepcopy(event)]
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):service.analyze(sid)
    return sid


def test_risk_requires_explicit_acceptance(service):
    with pytest.raises(ExportServiceError,match='xác nhận'):service.create_session(9,2026)
    with service.db._get_connection() as c:assert c.execute('SELECT count(*) FROM export_sessions').fetchone()[0]==0


def test_snapshot_and_idempotency(service):
    first=service.create_session(9,2026,True,'request1')
    again=service.create_session(9,2026,True,'request1')
    other_click=service.create_session(9,2026,True,'request2')
    assert first['session_id']==again['session_id']==other_click['session_id']
    session=service._session(first['session_id'])
    assert session['snapshot_data'][0]['content_hash']
    assert 'api_key' not in session['ai_settings']


@pytest.mark.parametrize('month,year,start,end',[(9,2026,'2026-08-15','2026-09-30'),(1,2026,'2025-12-15','2026-01-31'),(2,2028,'2028-01-15','2028-02-29')])
def test_period_timezone(service,month,year,start,end):
    a,b=service._get_date_range(month,year)
    assert a.isoformat().startswith(start) and b.isoformat().startswith(end)
    assert a.utcoffset().total_seconds()==25200 and b.hour==23 and b.second==59


def test_early_leave_alias_counts_like_reference(tmp_path):
    from src.utils.excel_export import export_to_excel
    path=tmp_path/'alias.xlsx'
    export_to_excel([{'person_id':'p','employee_name':'Invented A','event_date':'2026-09-05',
                     'event_type':'Nghỉ sớm','minutes_not_compensated':30,'source_message_ids':[]}],
                    {'month':9,'year':2026,'snapshot':[]},path)
    workbook=openpyxl.load_workbook(path)
    assert workbook['2_CHAM_CONG']['R5'].value==1


def test_legacy_sessions_readable_but_blocked(service):
    service.session_repo.create_session('old',9,2026)
    assert service.session_repo.get_session('old')
    with pytest.raises(ExportServiceError,match='snapshot'):service.export_to_excel('old')


def test_source_edit_blocks_export(service,event):
    sid=analyzed(service,event)
    service.message_repo.save_message('m1','test-room',1,'Người giả lập A','Đã sửa: OT 3h','2026-09-05T18:00:00+07:00')
    assert not service.verify_source_integrity(sid)['valid']
    with pytest.raises(ExportServiceError,match='nguồn đã đổi'):service.export_to_excel(sid)


def test_source_addition_blocks_export(service,event):
    sid=analyzed(service,event)
    service.message_repo.save_message('m2','test-room',1,'A','Tin mới','2026-09-06T08:00:00+07:00')
    assert not service.verify_source_integrity(sid)['valid']


def test_room_change_keeps_snapshot_room(service,event):
    sid=analyzed(service,event)
    service.config_repo.set('chatwork_room_id','another-room')
    assert service.verify_source_integrity(sid)['valid']


def test_pending_blocks_then_confirmation_exports(service,event):
    event['needs_review']=True;event['review_reason']='Cần đối chiếu'
    sid=analyzed(service,event)
    with pytest.raises(ExportServiceError,match='chưa hợp lệ'):service.export_to_excel(sid)
    row=service.get_results(sid)[0]
    service.review(sid,row['result_id'],{'reason':'Đã kiểm tra'})
    updated=service.get_results(sid)[0]
    assert updated['review_state']=='confirmed'
    assert updated['ai_raw_output']['reason']=='Công việc giả lập'
    assert service.session_repo.get_corrections(sid)
    result=service.export_to_excel(sid)
    assert Path(result['file_path']).exists()
    wb=openpyxl.load_workbook(result['file_path'])
    assert wb.sheetnames==['1_TIN_NHAN_GOC','2_CHAM_CONG','3_TONG_KET','4_CAN_KIEM_TRA']
    assert wb['2_CHAM_CONG']['A4'].value=='Event ID'
    assert wb['2_CHAM_CONG']['T5'].value==120
    assert wb['2_CHAM_CONG']['W5'].value=='=V5/60'
    assert 'BẢN TẠM' in wb['2_CHAM_CONG']['A1'].value
    assert wb['1_TIN_NHAN_GOC']['E5'].value.startswith('Người giả lập')
    wb.close()


def test_mandatory_error_cannot_be_confirmed_but_can_be_excluded(service,event):
    event['event_date']='2026-02-30'
    sid=analyzed(service,event);row=service.get_results(sid)[0]
    with pytest.raises(ExportServiceError,match='ngày không hợp lệ'):service.review(sid,row['result_id'])
    service.exclude_result(row['result_id'],sid)
    assert service.get_results(sid)==[]
    assert service.export_to_excel(sid)['event_count']==0


def test_outside_month_auto_excluded_but_september_message_for_october_kept(service,event):
    message=service.message_repo.get_by_id('m1')
    message.update(send_time='2026-09-23T08:00:00+07:00',body='Người giả lập A xin về sớm 30 phút ngày 23/09. Xin nghỉ ngày 01/10.')
    service.message_repo.save_messages_batch([message])
    sid=service.create_session(10,2026,True)['session_id']
    outside={**event,'event_date':'2026-09-23','event_type':'Về sớm','ot_minutes':0,'minutes_not_compensated':30}
    inside={**event,'event_date':'2026-10-01','event_type':'Nghỉ','ot_minutes':0,'event_subtype':'Phép năm','leave_days':1}
    adapter=Mock();adapter.analyze_attendance.return_value=[outside,inside]
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):
        assert service.analyze(sid)['status']=='ready'
    assert service.get_results_for_review(sid)==[]
    assert [r['event_date'] for r in service.get_results(sid)]==['2026-10-01']
    audit=service.get_results(sid,True)
    assert audit[0]['is_excluded'] and audit[0]['ai_raw_output']['event_date']=='2026-09-23'
    path=service.export_to_excel(sid)['file_path']
    w=openpyxl.load_workbook(path)
    try:
        assert w['2_CHAM_CONG']['E5'].value.month==10
        assert w['2_CHAM_CONG']['A6'].value is None
        assert w['1_TIN_NHAN_GOC']['A5'].value=='m1'
        assert w['4_CAN_KIEM_TRA']['F5'].value is None
    finally:w.close()


@pytest.mark.parametrize('changes',[{'source_message_ids':['fake']},{'ot_minutes':-1},{'event_date':'2026-02-30'},{'start_time':'18:00','end_time':'18:30'},{'confidence':'high'}])
def test_invalid_ai_is_blocked(service,event,changes):
    event.update(changes);sid=analyzed(service,event)
    assert service.get_results(sid)[0]['validation_errors']
    with pytest.raises(ExportServiceError):service.export_to_excel(sid)


def test_bad_quote_repaired_from_source_with_original_ai_preserved(service,event):
    event['identity_evidence']='Not an actual source quotation'
    sid=analyzed(service,event)
    row=service.get_results(sid)[0]
    assert not row['validation_errors']
    assert row['identity_evidence']=='Người giả lập A'
    assert row['ai_raw_output']['identity_evidence']=='Not an actual source quotation'


def test_single_clock_boundary_does_not_mean_invalid_clock(service,event):
    event['end_time']='11:30'
    sid=analyzed(service,event)
    assert not any('giờ bắt đầu/kết thúc không hợp lệ' in e for e in service.get_results(sid)[0]['validation_errors'])


def test_ai_timeout_keeps_failed_session(service):
    sid=service.create_session(9,2026,True)['session_id']
    with patch('src.services.export_service.create_ai_adapter',side_effect=TimeoutError('fake-key-never-log')):
        with pytest.raises(ExportServiceError):service.analyze(sid)
    assert service.session_repo.get_session(sid)['status']=='failed'
    assert 'fake-key-never-log' not in str(service.action_repo.get_events())


def test_ai_http_failure_saved_and_explained_without_provider_body(service):
    from src.adapters.ai_adapters import AIHTTPError
    sid=service.create_session(9,2026,True)['session_id']
    with patch('src.services.export_service.create_ai_adapter',side_effect=AIHTTPError(404)):
        with pytest.raises(ExportServiceError,match='Model không khả dụng'):
            service.analyze(sid)
    session=service.session_repo.get_session(sid)
    assert session['status']=='failed'
    assert 'HTTP 404' in session['error_message']


def test_ai_windows_key_failure_has_actionable_error(service):
    from src.utils.security import SecretError
    sid=service.create_session(9,2026,True)['session_id']
    with patch.object(service.config_repo,'get_ai_provider',side_effect=SecretError('do-not-log-secret')):
        with pytest.raises(ExportServiceError,match='Nhập lại API key'):
            service.analyze(sid)
    assert 'do-not-log-secret' not in service.session_repo.get_session(sid)['error_message']


def test_duplicate_analyze_does_not_call_ai(service,event):
    sid=analyzed(service,event)
    with patch('src.services.export_service.create_ai_adapter') as factory:
        assert service.analyze(sid)['existing']
        factory.assert_not_called()


def test_multiple_people_and_no_employee_catalog(service,event):
    sid=service.create_session(9,2026,True)['session_id']
    second={**event,'person_id':'P2','employee_name':'Người giả lập B','identity_evidence':'Người giả lập B'}
    with patch('src.services.export_service.create_ai_adapter') as factory:
        factory.return_value.analyze_attendance.return_value=[event,second]
        service.analyze(sid)
    assert len(service.get_results(sid))==2
    assert service.export_to_excel(sid)['event_count']==2
    with service.db._get_connection() as c:
        assert c.execute("SELECT count(*) FROM sqlite_master WHERE name='employees'").fetchone()[0]==0


def test_same_name_different_identities_never_merged(service,event):
    other=service.message_repo.get_by_id('m1')
    other.update(message_id='namesake',account_id=2)
    service.message_repo.save_messages_batch([other])
    sid=service.create_session(9,2026,True)['session_id']
    second={**event,'person_id':'P2','source_message_ids':['namesake']}
    with patch('src.services.export_service.create_ai_adapter') as factory:
        factory.return_value.analyze_attendance.return_value=[event,second];service.analyze(sid)
    assert len(service.get_results_for_review(sid))==2
    for row in service.get_results(sid):service.review(sid,row['result_id'])
    result=service.export_to_excel(sid)
    wb=openpyxl.load_workbook(result['file_path'])
    assert wb['3_TONG_KET']['A5'].value!=wb['3_TONG_KET']['A6'].value
    wb.close()


def test_multi_batch_registry_and_inconsistent_identity_blocked(service,event):
    service.message_repo.save_messages_batch([{'message_id':f'b{i}','room_id':'test-room','account_id':1,
        'name':event['employee_name'],'body':'Tin giả lập '+str(i),'send_time':'2026-09-05T18:00:00+07:00'} for i in range(41)])
    sid=service.create_session(9,2026,True)['session_id']
    seen=[]
    adapter=Mock()
    def analyze(batch,*args):
        seen.append(copy.deepcopy(adapter.identities) if isinstance(adapter.identities,list) else [])
        return [copy.deepcopy(event) if len(seen)==1 else {**event,'employee_name':'Tên khác','event_date':'2026-09-06'}]
    adapter.identities=[];adapter.analyze_attendance.side_effect=analyze
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):service.analyze(sid)
    assert len(seen)>=2 and seen[1][0]['person_id']=='P1'
    assert service.get_results_for_review(sid)
    assert any(r['validation_warnings'] or r['validation_errors'] for r in service.get_results(sid))
    with pytest.raises(ExportServiceError):service.export_to_excel(sid)


def test_timeout_splits_batch_without_losing_sources(service,event):
    from src.adapters.ai_adapters import AIResponseError
    second=service.message_repo.get_by_id('m1')
    second['message_id']='m2'
    service.message_repo.save_messages_batch([second])
    sid=service.create_session(9,2026,True)['session_id']
    calls=[]
    adapter=Mock()
    def analyze(batch,*args):
        calls.append([m['message_id'] for m in batch])
        if len(batch)>1:raise AIResponseError('timeout')
        return [{**event,'source_message_ids':[batch[0]['message_id']]}]
    adapter.analyze_attendance.side_effect=analyze
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):
        assert service.analyze(sid)['success']
    assert calls==[['m1','m2'],['m1'],['m2']]
    assert {r['source_message_ids'][0] for r in service.get_results(sid)}=={'m1','m2'}


def test_single_message_timeout_saved_without_partial_results(service):
    from src.adapters.ai_adapters import AIResponseError
    sid=service.create_session(9,2026,True)['session_id']
    adapter=Mock();adapter.analyze_attendance.side_effect=AIResponseError('timeout')
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):
        with pytest.raises(ExportServiceError,match='thời gian chờ'):service.analyze(sid)
    assert service.get_results(sid)==[]
    assert 'thời gian chờ' in service.session_repo.get_session(sid)['error_message']


def test_failed_batch_resumes_checkpoint_without_reanalyzing_completed_message(service,event):
    from src.adapters.ai_adapters import AIResponseError,AIHTTPError
    message=service.message_repo.get_by_id('m1');message['message_id']='m2'
    service.message_repo.save_messages_batch([message])
    sid=service.create_session(9,2026,True,'first')['session_id']
    adapter=Mock();adapter.analyze_attendance.side_effect=[AIResponseError('timeout'),[copy.deepcopy(event)],AIHTTPError(503)]
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):
        with pytest.raises(ExportServiceError,match='503'):service.analyze(sid)
    assert service.get_results(sid)==[]
    checkpoint=json.loads(service.session_repo.get_session(sid)['analysis_checkpoint'])
    assert checkpoint['completed']==['m1']
    assert service.create_session(9,2026,True,'retry-click')['session_id']==sid
    resumed=Mock();resumed.analyze_attendance.return_value=[{**event,'event_date':'2026-09-06','source_message_ids':['m2']}]
    with patch('src.services.export_service.create_ai_adapter',return_value=resumed):
        assert service.analyze(sid)['success']
    assert [m['message_id'] for m in resumed.analyze_attendance.call_args.args[0]]==['m2']
    assert len(service.get_results(sid))==2
    assert service.session_repo.get_session(sid)['analysis_checkpoint'] is None
    assert service.session_repo.get_session(sid)['error_message'] is None


def test_october_export_keeps_september_context_without_counting_it(service,event):
    message=service.message_repo.get_by_id('m1')
    message['send_time']='2026-09-25T18:00:00+07:00'
    service.message_repo.save_messages_batch([message])
    message.update(message_id='oct',send_time='2026-10-05T18:00:00+07:00')
    service.message_repo.save_messages_batch([message])
    sid=service.create_session(10,2026,True,'october')['session_id']
    adapter=Mock();adapter.analyze_attendance.return_value=[{**event,'event_date':'2026-10-05','source_message_ids':['oct']}]
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):
        assert service.analyze(sid)['status']=='ready'
    file=service.export_to_excel(sid)['file_path']
    workbook=openpyxl.load_workbook(file)
    try:
        assert workbook['2_CHAM_CONG']['T5'].value==120
        assert workbook['2_CHAM_CONG']['A6'].value is None
        assert workbook['1_TIN_NHAN_GOC']['A5'].value=='m1'
        assert workbook['1_TIN_NHAN_GOC']['A6'].value=='oct'
    finally:workbook.close()


@pytest.mark.parametrize('field,value',[('person_id',[]),('event_type',{}),('event_date',123),('employee_name',None),('source_message_ids',123)])
def test_malformed_identity_types_remain_reviewable(service,event,field,value):
    event[field]=value
    sid=analyzed(service,event)
    assert service.get_results(sid)[0]['review_state']=='pending'
    service.review(sid,service.get_results(sid)[0]['result_id'],exclude=True)
    assert Path(service.export_to_excel(sid)['file_path']).is_file()


def test_long_source_preserved_across_excel_cells(service,event):
    message=service.message_repo.get_by_id('m1')
    message['body']+=' x'*17000
    service.message_repo.save_messages_batch([message])
    sid=analyzed(service,event)
    path=service.export_to_excel(sid)['file_path']
    wb=openpyxl.load_workbook(path)
    try:
        assert ''.join(wb['1_TIN_NHAN_GOC'].cell(i,5).value for i in (5,6))==message['body']
    finally:wb.close()


def test_excluded_row_remains_in_audit_sheet(service,event):
    event['source_message_ids']=['fake'];sid=analyzed(service,event)
    service.review(sid,service.get_results(sid)[0]['result_id'],exclude=True)
    wb=openpyxl.load_workbook(service.export_to_excel(sid)['file_path'])
    try:assert wb['4_CAN_KIEM_TRA']['F5'].value=='Đã loại'
    finally:wb.close()


def test_snapshot_hash_tampering_blocks_export(service,event):
    sid=analyzed(service,event)
    snapshot=service._session(sid)['snapshot_data'];snapshot[0]['content_hash']='tampered'
    with service.db._get_connection() as c:c.execute('UPDATE export_sessions SET snapshot=? WHERE session_id=?',(json.dumps(snapshot),sid))
    with pytest.raises(ExportServiceError):service.export_to_excel(sid)


def test_portable_move_resolves_archived_file(service,event,tmp_path):
    sid=analyzed(service,event);old=Path(service.export_to_excel(sid)['file_path'])
    destination=tmp_path/'moved portable';destination.mkdir()
    shutil.move(str(Path(service.app_path)/'outputs'),str(destination/'outputs'))
    service.app_path=str(destination)
    resolved=service.session_file(service.session_repo.get_session(sid))
    assert resolved.is_file() and not old.exists()


def test_ot_used_once_and_cross_month_credit_not_deducted(service,event):
    ot={**event,'ot_used_minutes':30}
    leave={**event,'event_type':'Nghỉ bù OT','event_date':'2026-09-06','ot_minutes':0,
           'ot_used_minutes':30,'ot_credit_date':'2026-09-05','minutes_compensated':30}
    previous={**leave,'event_date':'2026-09-07','ot_credit_date':'2026-08-20','ot_used_minutes':60,'minutes_compensated':60}
    sid=service.create_session(9,2026,True)['session_id']
    with patch('src.services.export_service.create_ai_adapter') as factory:
        factory.return_value.analyze_attendance.return_value=[ot,leave,previous];service.analyze(sid)
    wb=openpyxl.load_workbook(service.export_to_excel(sid)['file_path'])
    try:
        sheet=wb['2_CHAM_CONG']
        assert sheet['T5'].value==120 and sheet['V5'].value==90
        assert sheet['U5'].value==0 and sheet['U6'].value==30 and sheet['U7'].value==60
    finally:wb.close()


def test_alias_and_proxy_sender_require_evidence(service,event):
    event['employee_name']='Tên suy đoán';event['identity_uncertain']=True
    sid=analyzed(service,event)
    assert service.get_results(sid)[0]['review_state']=='pending'


def test_duplicate_event_blocked(service,event):
    snapshot=service.message_repo.get_by_room('test-room')
    rows=validate_set([event,copy.deepcopy(event)],snapshot,9,2026)
    assert any('trùng' in e for e in rows[1]['validation_errors'])


def test_locked_monthly_file_preserves_previous(service,event):
    sid=analyzed(service,event)
    # Patched fixture resolves Downloads from pytest tmp_path, not app_path.
    from src.services.export_service import user_downloads_dir
    monthly=user_downloads_dir()/'Cham_cong_09_2026.xlsx';monthly.parent.mkdir();monthly.write_bytes(b'old-file')
    with pytest.raises(ExportServiceError,match='ghi đè'):service.export_to_excel(sid)
    import os
    original=os.replace
    def locked(a,b):
        if Path(b)==monthly:raise PermissionError('locked')
        return original(a,b)
    with patch('src.services.export_service.os.replace',side_effect=locked):
        with pytest.raises(ExportServiceError):service.export_to_excel(sid,True)
    assert monthly.read_bytes()==b'old-file'
    assert service.session_repo.get_session(sid)['status']=='ready'


def test_interrupted_analysis_recovered(service):
    sid=service.create_session(9,2026,True)['session_id'];service.session_repo.update_session(sid,status='analyzing')
    service.recover()
    assert service.session_repo.get_session(sid)['status']=='interrupted'


def test_template_preserves_layout(service,event):
    sid=analyzed(service,event);result=service.export_to_excel(sid)
    a=openpyxl.load_workbook(TEMPLATE);b=openpyxl.load_workbook(result['file_path'])
    for x,y in zip(a,b):
        assert list(x.merged_cells.ranges)==list(y.merged_cells.ranges)
        assert x.column_dimensions['A'].width==y.column_dimensions['A'].width
        assert x['A4']._style==y['A4']._style
        assert x.page_setup.orientation==y.page_setup.orientation
    a.close();b.close()


def test_export_needs_no_external_template_and_opens_summary(service,event):
    service.config_repo.set('excel_template', str(Path(service.app_path)/'missing-user-file.xlsx'))
    (Path(service.app_path)/'templates'/'template.xlsx').unlink()
    sid=analyzed(service,event)
    result=service.export_to_excel(sid)
    workbook=openpyxl.load_workbook(result['file_path'])
    try:
        assert workbook.active.title == '3_TONG_KET'
        assert workbook['3_TONG_KET'].freeze_panes == 'B5'
        assert workbook['3_TONG_KET'].auto_filter.ref == 'A4:M5'
        assert workbook['2_CHAM_CONG'].auto_filter.ref == 'A4:Z5'
        assert workbook['3_TONG_KET']['A5'].value == event['employee_name']
    finally:
        workbook.close()
