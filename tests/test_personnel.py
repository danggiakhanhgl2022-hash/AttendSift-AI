"""Personnel acceptance tests: invented people only, never the private roster."""
import copy
import json
from contextlib import contextmanager
from datetime import date
from io import BytesIO
from unittest.mock import Mock, patch
import openpyxl
import pytest
from src.services.employee_registry import EmployeeRegistry, KEY
from src.services.personnel_policy import classify, leave_claims
from cloud.app import finalize
from src.utils.excel_export import export_to_excel, verify_workbook
from tests.test_export_service import service
from tests.test_cloud_connections import SettingsRepository
from cloud.security import PASSWORD_HASHER


@pytest.fixture
def registry():
    data = {}
    return EmployeeRegistry(lambda key, default=None: copy.deepcopy(data.get(key, default)),
                            lambda key, value: data.update({key: copy.deepcopy(value)}))


def message(body, name='Người báo hộ'):
    return dict(message_id='m1', room_id='test-room', account_id=1,
                name=name, body=body, send_time='2026-10-05T18:00:00+07:00')


def roster(registry):
    registry.import_entries(['Lan Demo', 'Người Thử An', 'Người Mẫu An'])
    return registry.snapshot()


def event(name):
    return dict(employee_name=name, identity_evidence=name, source_message_ids=['m1'],
                event_type='Tăng ca', event_date='2026-10-05', person_id='source:'+name,
                ot_minutes=60, needs_review=False)


def test_registry_import_idempotency_and_soft_delete(registry):
    assert registry.import_entries(['Lan Demo', 'Lan Demo'])['added'] == 1
    employee = registry.read()['employees'][0]
    registry.upsert({'employee_code':'001','aliases':['Lan']}, employee['id'])
    before = registry.snapshot()
    registry.upsert({'active':False}, employee['id'])
    assert registry.import_entries(['Lan Demo'])['added'] == 0
    assert registry.read()['employees'][0]['active'] is False
    assert before['employees'][0]['active'] is True
    assert before['employees'][0]['aliases'] == ['Lan']
    assert registry.read()['employees'][0]['id'] == employee['id']


def test_unique_code_and_alias(registry):
    first=registry.upsert({'full_name':'Lan Demo','employee_code':'001','aliases':['Lan']})
    with pytest.raises(ValueError): registry.upsert({'full_name':'An Demo','employee_code':'001'})
    with pytest.raises(ValueError): registry.upsert({'full_name':'An Demo','aliases':['LAN']})
    registry.upsert({'full_name':'An Demo'})
    assert len(registry.read()['employees']) == 2
    with pytest.raises(ValueError): registry.upsert({'full_name':'x','active':'false'})


def test_empty_requires_explicit_confirmation(registry):
    with pytest.raises(ValueError): registry.snapshot()
    assert registry.snapshot(True)['employees'] == []
    with pytest.raises(ValueError): registry.snapshot('true')


@pytest.mark.parametrize('name,expected',[
    ('Lan Demo','official'), ('An Demo','contractor'),
    ('An Demo (CTV)','contractor'), ('An','unresolved'),
    ('Lan Demo (CTV)','unresolved'), ('','unresolved')])
def test_classification(registry,name,expected):
    row=event(name)
    assert classify(row,roster(registry),[message(name+' OT 1h')]) == expected


def test_alias_unicode_and_review_override(registry):
    official=registry.upsert({'full_name':'Lân Demo','aliases':['Lân']})
    row=event('La\u0302n')
    assert classify(row,registry.snapshot(),[message('Lân OT 1h')]) == 'official'
    assert row['employee_name']=='Lân Demo'
    ambiguous=event('Demo')
    classify(ambiguous,registry.snapshot(),[message('Demo OT 1h')])
    classify(ambiguous,registry.snapshot(),[message('Demo OT 1h')],
             {'classification':'official','employee_id':official['id']})
    assert ambiguous['classification_confirmed']
    with pytest.raises(ValueError):
        classify(event('X'),registry.snapshot(),[],{'classification':'official','employee_id':'unknown'})


@pytest.mark.parametrize('text,expected',[
 ('Em xin phép nghỉ ngày 05/10 và sáng ngày 06/10',{'2026-10-05':1,'2026-10-06':.5}),
 ('Em xin phép nghỉ sáng ngày 05/10',{'2026-10-05':.5}),
 ('Em xin phép nghỉ chiều ngày 05/10',{'2026-10-05':.5}),
 ('Em xin phép nghỉ sáng và chiều ngày 05/10',{'2026-10-05':1}),
 ('Em xin phép nghỉ sáng ngày 05/10 và chiều ngày 05/10',{'2026-10-05':1}),
 ('Em xin phép nghỉ sáng ngày 05/10 và sáng ngày 05/10',{'2026-10-05':.5}),
 ('Em xin phép nghỉ ngày 5-7/10',{'2026-10-05':1,'2026-10-06':1,'2026-10-07':1}),
 ('Em xin phép nghỉ 2 giờ ngày 05/10',{}),
])
def test_leave_durations(text,expected):
    actual=leave_claims(text,date(2026,10,1))
    assert {day:value['leave_days'] for day,value in actual.items()}==expected


def test_mixed_report_five_sheets_and_safe_raw(registry):
    staff=roster(registry)
    source=[message('Team Demo: ngày 5/10\nLan Demo OT 1h\nAn Demo (CTV) OT 2h')]
    results=finalize([],source,10,2026,'test',employee_registry=staff)
    active=[e for e in results if not e.get('is_excluded')]
    assert {(e['employee_name'],e['classification'],e['ot_minutes']) for e in active} == {
        ('Lan Demo','official',60),('An Demo (CTV)','contractor',120)}
    assert all(not e['validation_errors'] for e in active)
    buffer=BytesIO()
    export_to_excel(results,{'month':10,'year':2026,'snapshot':source,
                            'ai_settings':{'employee_registry':staff}},buffer)
    buffer.seek(0);assert verify_workbook(buffer)
    buffer.seek(0);workbook=openpyxl.load_workbook(buffer,data_only=True)
    try:
        assert workbook.sheetnames[-1]=='5_CTV'
        assert workbook['2_CHAM_CONG']['F5'].value=='Lan Demo'
        assert workbook['3_TONG_KET']['I5'].value==60
        assert workbook['5_CTV']['E5'].value==120
        assert workbook['5_CTV']['J5'].value==source[0]['body']
    finally:workbook.close()


def test_generic_leave_has_two_rows_and_no_missing_days_warning(registry):
    registry.import_entries(['Lan Demo'])
    source=[message('Em xin phép nghỉ ngày 05/10 (thứ 2) và sáng ngày 06/10 (thứ 3)\nLý do: Việc giả lập','Lan Demo')]
    rows=finalize([],source,10,2026,'test',employee_registry=registry.snapshot())
    active=[e for e in rows if not e.get('is_excluded')]
    assert [(e['event_date'],e['leave_days']) for e in active]==[('2026-10-05',1),('2026-10-06',.5)]
    assert all(not e['validation_errors'] for e in active)
    assert all(not e.get('event_subtype') for e in active)


def test_local_snapshot_and_reclassification_preserve_history(service, registry):
    service.employee_registry=registry
    registry.import_entries(['Lan Demo', 'Bình Demo'])
    source=message('Team Demo: ngày 5/10\nLan Demo OT 1h\nAn Demo (CTV) OT 2h')
    service.message_repo.save_messages_batch([source])
    sid=service.create_session(10,2026,True)['session_id']
    snapshot=json.loads(service._session(sid)['ai_settings'])['employee_registry']
    registry.upsert({'active':False},registry.read()['employees'][0]['id'])
    adapter=Mock();adapter.analyze_attendance.return_value=[]
    with patch('src.services.export_service.create_ai_adapter',return_value=adapter):
        service.analyze(sid)
    rows=service.get_results(sid)
    assert next(e for e in rows if e['employee_name']=='Lan Demo')['classification']=='official'
    assert json.loads(service._session(sid)['ai_settings'])['employee_registry']==snapshot
    contractor=next(e for e in rows if e['classification']=='contractor')
    service.review(sid,contractor['result_id'],{'classification':'official','employee_id':snapshot['employees'][1]['id']})
    assert all(e['classification']=='official' for e in service.get_results(sid))


def test_local_employee_endpoints_csrf_and_persistence(tmp_path):
    from src.app.server import create_app
    app=create_app(str(tmp_path));client=app.test_client()
    import re
    token=re.search(r'name="csrf-token" content="([^"]+)"',client.get('/').get_data(as_text=True))[1]
    assert client.post('/api/employees',json={'full_name':'Lan Demo'}).status_code==403
    headers={'X-CSRF-Token':token}
    response=client.post('/api/employees',json={'full_name':'Lan Demo'},headers=headers)
    assert response.status_code==200
    employee=response.json['data']
    assert client.post('/api/employees/'+employee['id'],json={'active':False},headers=headers).status_code==200
    again=create_app(str(tmp_path)).test_client()
    assert again.get('/api/employees').json['data']['employees'][0]['active'] is False


def test_cloud_employee_endpoints_require_login_and_csrf():
    class Repo(SettingsRepository):
        @contextmanager
        def lease(self,*args,**kwargs):yield 'owner'
    from cloud.app import create_app
    repo=Repo()
    app=create_app(repository=repo,config={
        'CHAMCONG_SESSION_SECRET':'synthetic-private-session-secret-1234567890',
        'CHAMCONG_ADMIN_PASSWORD_HASH':PASSWORD_HASHER.hash('synthetic-password'),
        'CHAMCONG_ALLOWED_HOSTS':'localhost','CHAMCONG_INSECURE_LOCAL':'1'})
    client=app.test_client()
    assert client.get('/api/employees').status_code==401
    login=client.post('/login',json={'password':'synthetic-password'})
    token=login.json['data']['csrf_token']
    assert client.post('/api/employees',json={'full_name':'Lan Demo'}).status_code==403
    response=client.post('/api/employees/import',json={'employees':['Lan Demo']},headers={'X-CSRF-Token':token})
    assert response.status_code==200
    assert KEY in repo.settings
    assert client.get('/api/employees').json['data']['employees'][0]['full_name']=='Lan Demo'


def test_registry_seed_does_not_restore_deleted_employee(registry,tmp_path):
    path=tmp_path/'seed.json'
    path.write_text(json.dumps({'employees':['Lan Demo']}),encoding='utf-8')
    registry.seed_once(path)
    item=registry.read()['employees'][0]
    registry.upsert({'active':False},item['id'])
    registry.seed_once(path)
    assert registry.read()['employees'][0]['active'] is False
    assert len(registry.read()['employees'])==1


def test_batch_import_validation_is_atomic(registry):
    with pytest.raises(ValueError):
        registry.import_entries([{'full_name':'Lan Demo'}, {'full_name':''}])
    assert registry.read()['employees']==[]


def test_ai_cannot_invent_official_full_name(registry):
    registry.import_entries(['Lan Demo'])
    row=event('Lan Demo')
    row['identity_evidence']='An'
    classify(row,registry.snapshot(),[message('An OT 1h')])
    assert row['classification']=='unresolved'


def test_unrelated_ctv_word_does_not_change_first_person_leave(registry):
    registry.import_entries(['Lan Demo'])
    source=[message('Em xin phép nghỉ sáng ngày 05/10\nEm nhờ CTV hỗ trợ công việc.','Lan Demo')]
    active=[e for e in finalize([],source,10,2026,'test',employee_registry=registry.snapshot()) if not e.get('is_excluded')]
    assert len(active)==1
    assert active[0]['classification']=='official'
    assert active[0]['leave_days']==.5


@pytest.mark.parametrize('duration',['2h','30ph','2 giờ','30 phút'])
def test_hour_based_leave_does_not_become_full_day(duration):
    assert leave_claims('Em xin phép nghỉ '+duration+' ngày 05/10',date(2026,10,1))=={}


def test_contractor_long_source_and_formula_safety(registry):
    staff=roster(registry)
    raw='=HYPERLINK("https://invalid.example")\n'+'x'*65000
    row={**event('An Demo'),'classification':'contractor','classification_reason':'Synthetic',
         'event_id':'E1','review_state':'valid','ot_minutes':120}
    stream=BytesIO()
    export_to_excel([row],{'month':10,'year':2026,'snapshot':[message(raw)],
                          'ai_settings':{'employee_registry':staff}},stream)
    stream.seek(0)
    workbook=openpyxl.load_workbook(stream,data_only=False)
    try:
        sheet=workbook['5_CTV']
        assert ''.join(sheet.cell(i,10).value for i in (5,6,7))==raw
        assert sheet['J5'].data_type=='s'
        assert sum(sheet.cell(i,5).value for i in (5,6,7))==120
    finally:workbook.close()
