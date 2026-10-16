"""Exported reports can be reviewed without trusting stale validation flags."""
import copy
import json
from io import BytesIO
from unittest.mock import Mock, patch

import openpyxl

from tests.test_export_service import service, event, analyzed
from tests.test_personnel_cloud_flow import FakeRepo, make_app, login
from src.services.employee_registry import EmployeeRegistry


def test_local_export_revalidates_the_same_rows_written_to_excel(service, event):
    sid = analyzed(service, event)
    result = service.get_results(sid)[0]
    result.update(ot_minutes=-5, validation_errors=[], validation_warnings=[])
    with service.db._get_connection() as conn:
        conn.execute('UPDATE ai_results SET payload=? WHERE result_id=?',
                     (json.dumps(result), result['result_id']))
    workbook = openpyxl.load_workbook(service.export_to_excel(sid)['file_path'])
    try:
        assert workbook['2_CHAM_CONG']['A5'].value is None
        assert workbook['3_TONG_KET']['A5'].value is None
        assert workbook['4_CAN_KIEM_TRA']['F5'].value == 'Cần kiểm tra (chưa tính công)'
        assert 'ot_minutes' in workbook['4_CAN_KIEM_TRA']['E5'].value
    finally:
        workbook.close()


def test_local_review_after_export_preserves_paid_break_minutes(service):
    directory = {}
    service.employee_registry = EmployeeRegistry(
        lambda key, default=None: copy.deepcopy(directory.get(key, default)),
        lambda key, value: directory.update({key: copy.deepcopy(value)}))
    service.employee_registry.import_entries(['Nhân Viên Demo'])
    source = service.message_repo.get_by_id('m1')
    source.update(name='Nhân Viên Demo', body='Em xin phép về sớm 30 phút không làm bù.',
                  send_time='2026-10-14T16:30:00+07:00')
    service.message_repo.save_messages_batch([source])
    sid = service.create_session(10, 2026, True)['session_id']
    adapter = Mock()
    adapter.analyze_attendance.return_value = []
    with patch('src.services.export_service.create_ai_adapter', return_value=adapter):
        service.analyze(sid)
    row = service.get_results(sid)[0]
    assert row['paid_break_minutes'] == 30
    service.export_to_excel(sid)
    service.review(sid, row['result_id'], {'reason': 'Đã đối chiếu tin nguồn'})
    assert service.get_results(sid)[0]['paid_break_minutes'] == 30
    workbook = openpyxl.load_workbook(service.export_to_excel(sid, True)['file_path'], data_only=True)
    try:
        assert workbook['2_CHAM_CONG']['AB5'].value == 30
        assert workbook['2_CHAM_CONG']['O5'].value == 0
    finally:
        workbook.close()


def cloud_session():
    repo = FakeRepo()
    client = make_app(repo).test_client()
    headers = login(client)
    client.post('/api/employees/import', json={'employees': ['Lan Demo']}, headers=headers)
    sid = client.post('/api/sessions', json={'month': 10, 'year': 2026, 'accept_risk': True},
                      headers=headers).json['data']['session_id']
    assert client.post(f'/api/sessions/{sid}/analyze', json={}, headers=headers).status_code == 200
    return repo, client, headers, sid


def test_cloud_export_can_be_reviewed_and_exported_again():
    repo, client, headers, sid = cloud_session()
    row = next(e for e in repo.records[sid]['results'] if e.get('classification') == 'official')
    assert client.post(f'/api/sessions/{sid}/export', json={}, headers=headers).status_code == 200
    before = copy.deepcopy(repo.files[sid])
    reviewed = client.post(f'/api/sessions/{sid}/review',
                           json={'event_id': row['event_id'], 'changes': {'reason': 'Đã đối chiếu'}},
                           headers=headers)
    assert reviewed.status_code == 200, reviewed.json
    assert repo.files[sid] == before  # Existing download stays intact until regenerated.
    assert client.post(f'/api/sessions/{sid}/export', json={}, headers=headers).status_code == 200
    assert any(e['review_state'] == 'confirmed' for e in repo.records[sid]['results'])


def test_cloud_export_revalidates_stored_results_before_counting():
    repo, client, headers, sid = cloud_session()
    row = next(e for e in repo.records[sid]['results'] if e.get('classification') == 'official')
    row.update(ot_minutes=-5, validation_errors=[], validation_warnings=[], review_state='valid')
    assert client.post(f'/api/sessions/{sid}/export', json={}, headers=headers).status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(repo.files[sid]['content']), data_only=True)
    try:
        assert workbook['2_CHAM_CONG']['A5'].value is None
        assert workbook['3_TONG_KET']['A5'].value is None
        assert workbook['4_CAN_KIEM_TRA']['F5'].value == 'Cần kiểm tra (chưa tính công)'
    finally:
        workbook.close()
