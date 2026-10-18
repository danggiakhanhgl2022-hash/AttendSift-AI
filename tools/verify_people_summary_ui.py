"""Read Excel on overview, keep payroll intact and verify the single neural scene."""
from pathlib import Path
import sys,tempfile,threading,json,logging,io,hashlib
from PIL import Image,ImageChops
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from src.app.server import create_app
from werkzeug.serving import make_server
from playwright.sync_api import sync_playwright
import openpyxl

def make_book(ot=120):
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name in ['1_TIN_NHAN_GOC','2_CHAM_CONG','3_TONG_KET','4_CAN_KIEM_TRA','5_CTV']:wb.create_sheet(name)
    for i,(name,minutes) in enumerate([('Nhân viên Demo A',ot),('Nhân viên Demo B',60)],5):
        wb['2_CHAM_CONG'].cell(i,1,'EVT'+str(i));wb['2_CHAM_CONG'].cell(i,6,name)
        wb['3_TONG_KET'].cell(i,1,name);wb['3_TONG_KET'].cell(i,9,minutes);wb['3_TONG_KET'].cell(i,14,0)
    sheet=wb['5_CTV'];sheet.cell(10,1,'TỔNG THEO NGƯỜI')
    for i,(name,minutes) in enumerate([('Cầm',720),('Trắng',300),('Ny',450),('My (CTV)',30)],11):
        for col,value in enumerate(['DEMO-'+str(i),name,minutes,0],1):sheet.cell(i,col,value)
    wb['4_CAN_KIEM_TRA'].cell(5,6,'Cần kiểm tra (chưa tính công)')
    out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()

def main():
    logging.getLogger('werkzeug').setLevel(logging.ERROR);report={};images=ROOT/'docs/images';images.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='excel-overview-',ignore_cleanup_errors=True) as temp:
        app=create_app(temp);services=app.extensions['services'];sid='DEMO-EXCEL'
        services['sessions'].create_session(sid,10,2026)
        # Deliberately misleading API classification must not override the Excel sheets.
        services['sessions'].save_results(sid,[dict(event_id='WRONG-CLASS',source_id='DEMO-S0',employee_name='Nhân viên Demo A',event_type='Tăng ca',event_date='2026-10-16')])
        with services['db']._get_connection() as c:c.execute('UPDATE ai_results SET payload=?',(json.dumps({'classification':'contractor','ot_minutes':999}),));c.commit()
        path=Path(temp)/'outputs/sessions'/sid/'chamcong-demo.xlsx';path.parent.mkdir(parents=True);path.write_bytes(make_book())
        services['sessions'].update_session(sid,status='exported',file_path=str(path));before_hash=hashlib.sha256(path.read_bytes()).hexdigest()
        server=make_server('127.0.0.1',0,app,threaded=True);threading.Thread(target=server.serve_forever,daemon=True).start()
        try:
            with sync_playwright() as pw:
                browser=pw.chromium.launch(channel='msedge',headless=True,args=['--enable-unsafe-swiftshader'])
                page=browser.new_page(viewport={'width':1600,'height':1100},reduced_motion='reduce');errors=[];posts=[]
                page.on('pageerror',lambda e:errors.append(str(e)));page.on('request',lambda r:posts.append(r.url) if r.method!='GET' else None)
                page.goto(f'http://127.0.0.1:{server.server_port}',wait_until='networkidle')
                page.wait_for_function("() => document.querySelector('#people-ot')?.textContent==='180'")
                assert page.locator('#people-count').inner_text()=='2'
                assert 'Nhân viên chính thức' in page.locator('.people-summary h2').inner_text()
                assert '3_TONG_KET' in page.locator('#people-scope').inner_text()
                page.locator('[data-people-filter="contractor"]').click();assert page.locator('#people-ot').inner_text()=='1.500'
                assert 'Nhân viên Demo A' not in page.locator('#people-rows').inner_text()
                report['official_from_attendance_summary_ctv_from_ctv_sheet']=True
                page.locator('[data-people-filter="official"]').click()
                page.locator('#people-summary').screenshot(path=str(images/'people-summary.png'))
                upload=Path(temp)/'another-demo.xlsx';upload.write_bytes(make_book(240))
                page.locator('#people-workbook-file').set_input_files(upload)
                page.wait_for_function("() => document.querySelector('#people-ot').textContent==='300'")
                page.locator('#people-session-file').click();page.wait_for_function("() => document.querySelector('#people-ot').textContent==='180'")
                report['choose_file_and_return_to_session_file']=True
                assert page.get_by_text('02 / OPERATIONS ROOM',exact=True).count()==0
                assert page.evaluate('typeof window.createOperationsStudio')=='undefined'
                page.locator('#dash-brain-slot').scroll_into_view_if_needed();page.wait_for_function('() => CinematicScene.getState().frames>3')
                frames=page.evaluate('CinematicScene.getState().frames');page.wait_for_timeout(350);assert page.evaluate('CinematicScene.getState().frames')>frames
                a=Image.open(io.BytesIO(page.locator('.neural-panel').screenshot())).convert('RGB');page.wait_for_timeout(300)
                b=Image.open(io.BytesIO(page.locator('.neural-panel').screenshot())).convert('RGB');assert ImageChops.difference(a,b).getbbox()
                frames=[]
                for _ in range(8):frames.append(Image.open(io.BytesIO(page.locator('.neural-panel').screenshot())).convert('RGB'));page.wait_for_timeout(150)
                frames[0].save(images/'neural-motion.gif',save_all=True,append_images=frames[1:],duration=200,loop=0)
                page.evaluate('scrollTo(0,0)');page.screenshot(path=str(images/'neural-terminal.png'),full_page=True)
                for key in ['sync','export','history','settings','employees']:
                    page.locator(f'.nav [data-section="{key}"]').click();assert page.locator(f'#{key}-section').is_visible()
                    page.screenshot(path=str(images/f'ui-{key}.png'))
                page.locator('#dashboard-nav').click();page.set_viewport_size({'width':390,'height':844});page.wait_for_timeout(300)
                assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
                report['single_neural_scene_motion_original_tabs_mobile']=True
                assert not errors,errors;assert all(url.endswith('/api/dashboard/workbook-summary') for url in posts),posts
                report['javascript_errors']=errors;report['only_upload_parsing_post']=True;browser.close()
        finally:server.shutdown();server.server_close()
        assert hashlib.sha256(path.read_bytes()).hexdigest()==before_hash
        with services['db']._get_connection() as c:
            assert c.execute('select count(*) from ai_results').fetchone()[0]==1
            assert json.loads(c.execute('select payload from ai_results').fetchone()[0])['classification']=='contractor'
        report['workbook_and_original_classification_unchanged']=True
    out=ROOT/'docs/verification';out.mkdir(parents=True,exist_ok=True);(out/'workbook-overview.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report,indent=2))
if __name__=='__main__':main()
