"""Export one cleaned content column; source evidence stays in the database."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from .excel_export import safe_cell
from .message_cleaning import clean_message_body


def export_raw_messages(messages, output_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Tin nhan'
    sheet.append(['Mã tin nhắn', 'Thời gian gửi', 'Người gửi', 'Phòng',
                  'Nội dung tin nhắn'])
    for message in sorted(messages, key=lambda m: str(m['send_time'])):
        original = message.get('body') or ''
        cleaned = clean_message_body(original, message.get('name'))
        # Excel limits cell size: preserve all text using continuation rows.
        count = max(1, (len(cleaned) + 31999) // 32000)
        for part in range(count):
            offset = part * 32000
            values = [message['message_id'], message['send_time'], message.get('name'),
                      message.get('room_name') or message.get('room_id'),
                      cleaned[offset:offset + 32000]]
            row = sheet.max_row + 1
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row, col)
                safe_cell(cell, value)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for col, width in {'A':22, 'B':28, 'C':28, 'D':22, 'E':90}.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output_path)
    workbook.close()
