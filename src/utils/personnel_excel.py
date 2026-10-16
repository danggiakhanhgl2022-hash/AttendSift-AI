"""Write contractor details and totals without adding to official payroll."""
from collections import defaultdict
from openpyxl.styles import Alignment, Font, PatternFill
from .excel_export import safe_cell


def append_personnel_sheets(wb, official, contractors, source, month, year):
    official_sheet = wb['2_CHAM_CONG']
    official_sheet.cell(4, 27, 'Số ngày nghỉ (mọi loại)')
    official_sheet.column_dimensions['AA'].width = 24
    for index, event in enumerate(official, 5):
        official_sheet.cell(index, 27, event.get('leave_days') or 0)
    # General leave is separate from paid/unpaid leave columns: no invented payroll policy.
    totals = defaultdict(float)
    for event in official:
        totals[event.get('person_id', event.get('employee_name'))] += event.get('leave_days') or 0
    summary = wb['3_TONG_KET']
    summary.cell(4, 14, 'Tổng ngày nghỉ (mọi loại)')
    summary.column_dimensions['N'].width = 26
    for row in range(5, summary.max_row + 1):
        label = summary.cell(row, 1).value
        total = sum(value for identity, value in totals.items()
                    if any(e.get('person_id', e.get('employee_name')) == identity and
                           (e.get('employee_name') == label or f"{e.get('employee_name')} [{identity}]" == label)
                           for e in official))
        if label:
            summary.cell(row, 14, total)
    for sheet in (official_sheet, summary):
        last = sheet.max_row
        end = sheet.cell(4, sheet.max_column).column_letter
        sheet.auto_filter.ref = f'A4:{end}{last}'
        sheet.print_area = f'A1:{end}{last}'

    sheet = wb.create_sheet('5_CTV')
    safe_cell(sheet.cell(1, 1), f'CTV — THÁNG {month:02d}/{year}')
    headers = ['Event ID', 'Họ tên', 'Ngày sự kiện', 'Loại', 'OT (phút)',
               'Số ngày nghỉ', 'Căn cứ phân loại', 'Trạng thái', 'Source ID',
               'Tin nhắn gốc', 'Diễn giải']
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(4, column)
        safe_cell(cell, header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='28576E')
    row = 5
    totals = {}
    for event in contractors:
        ids = event.get('source_message_ids') or []
        raw = '\n\n'.join(source[i]['body'] for i in ids if i in source)
        chunks = [raw[i:i+32000] for i in range(0, len(raw), 32000)] or ['']
        ot = event.get('ot_minutes', 0) if event.get('event_type') in {'OT', 'Tăng ca'} else 0
        leave = event.get('leave_days') or 0
        identity = event.get('person_id') or event.get('employee_name')
        if identity not in totals:
            totals[identity] = [event.get('employee_name'), 0, 0]
        totals[identity][1] += ot
        totals[identity][2] += leave
        for part, text in enumerate(chunks):
            values = [event.get('event_id'), event.get('employee_name'), event.get('event_date'),
                      event.get('event_type') if part == 0 else 'Nối tin nguồn',
                      ot if part == 0 else 0, leave if part == 0 else 0,
                      event.get('classification_reason'), event.get('review_state'),
                      ', '.join(ids), text, event.get('reason')]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row, column)
                safe_cell(cell, value)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            sheet.row_dimensions[row].height = 58
            row += 1
    sheet.freeze_panes = 'C5'
    sheet.auto_filter.ref = f'A4:K{max(5,row-1)}'
    row += 2
    for column, value in enumerate(['TỔNG THEO NGƯỜI', 'Họ tên', 'Tổng OT (phút)', 'Tổng ngày nghỉ'], 1):
        safe_cell(sheet.cell(row, column), value)
    for identity, (name, ot, leave) in sorted(totals.items()):
        row += 1
        for column, value in enumerate([identity, name, ot, leave], 1):
            safe_cell(sheet.cell(row, column), value)
    for column, width in {'A':24,'B':28,'C':24,'D':22,'E':15,'F':17,'G':52,'H':18,'I':30,'J':90,'K':44}.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.tabColor = '7755AA'
