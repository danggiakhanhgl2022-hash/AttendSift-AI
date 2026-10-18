"""Read existing Excel sheets for overview display; never classify or write payroll."""
from collections import Counter
from io import BytesIO
import hashlib
from math import isfinite
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import openpyxl
from flask import jsonify, request


def read_workbook_summary(content, filename):
    if not str(filename).lower().endswith('.xlsx'):
        raise ValueError('Chọn file bảng chấm công .xlsx.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries=archive.infolist()
            if len(entries)>1500 or sum(e.file_size for e in entries)>100*1024*1024:
                raise ValueError('Workbook quá lớn để hiển thị tổng quan.')
        workbook=openpyxl.load_workbook(BytesIO(content),data_only=True,read_only=True,keep_links=False)
    except (BadZipFile, OSError, KeyError) as error:
        raise ValueError('File Excel không hợp lệ.') from error
    try:
        needed={'2_CHAM_CONG','3_TONG_KET','5_CTV'}
        if not needed.issubset(workbook.sheetnames):
            raise ValueError('File cần có sheet 2_CHAM_CONG, 3_TONG_KET và 5_CTV.')
        def rows(name):
            sheet=workbook[name]
            if sheet.max_row>30000 or sheet.max_column>100:
                raise ValueError('Workbook vượt giới hạn đọc tổng quan.')
            return sheet.iter_rows(values_only=True)
        def cell(row,index):return row[index-1] if len(row)>=index else None
        def number(value):
            if isinstance(value,(int,float)) and not isinstance(value,bool) and isfinite(value):return value
            return None
        def text(value):return str(value or '').strip()
        events=Counter();seen=set()
        for index,row in enumerate(rows('2_CHAM_CONG'),1):
            if index<5:continue
            identity=text(cell(row,6));event=text(cell(row,1))
            if identity and event and (identity,event) not in seen:events[identity]+=1;seen.add((identity,event))
        people=[]
        for index,row in enumerate(rows('3_TONG_KET'),1):
            if index<5:continue
            name=text(cell(row,1))
            if name:
                people.append(dict(id='—',name=name,classification='official',ot=number(cell(row,9)),leave=number(cell(row,14)),count=events[name],sheet='3_TONG_KET'))
        contractor_counts=Counter();in_totals=False
        for row in rows('5_CTV'):
            first=text(cell(row,1))
            if first.upper()=='TỔNG THEO NGƯỜI':in_totals=True;continue
            if in_totals:
                name=text(cell(row,2))
                if first and name:people.append(dict(id=first,name=name,classification='contractor',ot=number(cell(row,3)),leave=number(cell(row,4)),count=contractor_counts[name],sheet='5_CTV'))
            elif first.startswith('EVT') or first.startswith('EVENT'):
                contractor_counts[text(cell(row,2))]+=1
        pending=0
        if '4_CAN_KIEM_TRA' in workbook.sheetnames:
            for row in rows('4_CAN_KIEM_TRA'):
                if 'cần kiểm tra' in text(cell(row,6)).casefold():pending+=1
        return dict(available=True,filename=Path(filename).name,people=people,pending=pending,
                    source='excel',missing_cached_values=any(p['ot'] is None or p['leave'] is None for p in people))
    finally:workbook.close()


def register_workbook_view(app,sessions,export):
    @app.get('/api/export/<sid>/workbook-summary')
    def workbook_summary(sid):
        session=sessions.get_session(sid)
        if not session or session.get('status')!='exported' or not session.get('file_path'):
            return jsonify(success=True,data={'available':False})
        path=export.session_file(session)
        if not path:return jsonify(success=True,data={'available':False})
        content=path.read_bytes()
        if session.get('file_hash') and hashlib.sha256(content).hexdigest()!=session['file_hash']:
            raise ValueError('File đã thay đổi sau khi xuất; chọn bản Excel cần xem bằng nút Đọc file Excel.')
        return jsonify(success=True,data=read_workbook_summary(content,path.name))

    @app.post('/api/dashboard/workbook-summary')
    def uploaded_workbook_summary():
        file=request.files.get('file')
        if file is None:raise ValueError('Chọn file Excel để hiển thị.')
        # Kept in memory only; no database, source message or Excel file is changed.
        return jsonify(success=True,data=read_workbook_summary(file.read(),file.filename))
