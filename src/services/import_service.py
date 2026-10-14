"""Import historical source messages; never depend on employee directories."""
import csv
import io
import json
import hashlib
import re
from zipfile import ZipFile
from datetime import datetime
import pytz
from openpyxl import load_workbook

TZ = pytz.timezone('Asia/Ho_Chi_Minh')


class ImportService:
    def __init__(self, db, messages, config):
        self.db, self.messages, self.config = db, messages, config

    def import_file(self, filename, content):
        if filename.lower().endswith('.json'):
            data = json.loads(content.decode('utf-8-sig'))
            if isinstance(data, list):
                data = {'MessageLog': data}
            if not isinstance(data, dict) or 'nodes' in data:
                raise ValueError('Hãy nhập dữ liệu MessageLog, không phải file cấu hình workflow n8n.')
        elif filename.lower().endswith('.csv'):
            data = {'MessageLog': list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))}
        elif filename.lower().endswith('.xlsx'):
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('File Excel quá lớn sau giải nén (tối đa 40 MB).')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            data = {}
            try:
                for sheet in workbook:
                    if sheet.title == '1_TIN_NHAN_GOC':
                        room = self.config.get('chatwork_room_id')
                        if not room:
                            raise ValueError('Chọn phòng Chatwork trước khi nhập sheet tin nhắn gốc.')
                        records=[]
                        for row in sheet.iter_rows(min_row=5, values_only=True):
                            if not row or not row[0]:
                                continue
                            try:
                                day=row[1] if isinstance(row[1],datetime) else datetime.fromisoformat(str(row[1]))
                                hour,minute=map(int,str(row[2]).split(':')[:2])
                                sent=day.replace(hour=hour,minute=minute)
                            except (ValueError,TypeError,IndexError):
                                raise ValueError(f'Tin {row[0]}: ngày hoặc giờ không hợp lệ.') from None
                            body=str(row[4] or '')
                            digest=hashlib.sha256((str(room)+'|'+str(row[0])+'|'+sent.isoformat()+'|'+str(row[3])+'|'+body).encode()).hexdigest()
                            records.append({'message_id':'excel:'+digest,'room_id':room,'sender_name':row[3],
                                            'send_datetime':sent,'raw_body':body})
                        data['MessageLog']=records
                        continue
                    if sheet.title != 'MessageLog':
                        continue
                    rows = iter(sheet.values)
                    headers = [str(v or '').strip() for v in next(rows, [])]
                    data[sheet.title] = [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]
            finally:
                workbook.close()
        else:
            raise ValueError('Chỉ hỗ trợ XLSX, CSV hoặc JSON.')
        if 'MessageLog' not in data or not isinstance(data['MessageLog'],list):
            raise ValueError('Không tìm thấy MessageLog hoặc 1_TIN_NHAN_GOC. Không thể lấy tin nguồn từ bảng tổng kết.')
        prepared = []
        for index, row in enumerate(data.get('MessageLog', []), 1):
            if isinstance(row,dict) and isinstance(row.get('json'),dict):
                row=row['json']
            if not isinstance(row, dict):
                raise ValueError(f'MessageLog dòng {index}: dữ liệu không hợp lệ.')
            mid, room = str(row.get('message_id') or '').strip(), str(row.get('room_id') or '').strip()
            body = row.get('raw_body', row.get('body'))
            if not mid or not room or not isinstance(body, str):
                raise ValueError(f'MessageLog dòng {index}: thiếu message_id, room_id hoặc raw_body.')
            try:
                raw = row.get('send_datetime') or row.get('send_time')
                if isinstance(raw, datetime):
                    sent = raw
                elif raw:
                    sent = datetime.fromisoformat(re.sub(r'\s+', ' ', str(raw).strip()).replace('Z', '+00:00'))
                else:
                    number=str(row['send_time_unix']).strip()
                    if re.fullmatch(r'\d{1,3}(\.\d{3})+(,\d+)?',number):
                        number=number.replace('.','').replace(',','.')
                    sent = datetime.fromtimestamp(float(number), TZ)
                sent = TZ.localize(sent) if sent.tzinfo is None else sent.astimezone(TZ)
            except (ValueError, TypeError, KeyError, OverflowError):
                raise ValueError(f'MessageLog dòng {index}: thời gian gửi không hợp lệ.') from None
            prepared.append({'message_id': mid, 'room_id': room,
                             'account_id': row.get('sender_account_id', row.get('account_id')),
                             'name': str(row.get('sender_name', row.get('name')) or ''),
                             'body': body, 'send_time': sent.isoformat()})
        # All rows are validated before one atomic transaction. Never overwrite live edits.
        with self.db.operation_lock, self.db._get_connection() as conn:
            conn.execute('BEGIN IMMEDIATE')
            added = []
            skipped = 0
            seen = set()
            for message in prepared:
                if message['message_id'] in seen or conn.execute('SELECT 1 FROM messages WHERE message_id=?', (message['message_id'],)).fetchone():
                    skipped += 1
                else:
                    added.append(message)
                seen.add(message['message_id'])
            self.messages.save_messages_batch(added, conn)
        return {'added':len(added), 'skipped':skipped,
                'first_message':min((m['send_time'] for m in prepared),default=None),
                'last_message':max((m['send_time'] for m in prepared),default=None),
                'room_ids':sorted({m['room_id'] for m in prepared}),
                'imported_from_report':any(m['message_id'].startswith('excel:') for m in prepared)}
