"""Synthetic regression cases; no actual employee names, Chatwork IDs or messages."""
from io import BytesIO
import openpyxl

from src.services.personnel_policy import classify
from src.services.source_reconciliation import reconcile_explicit_records
from src.services.validation import validate_set
from src.utils.excel_export import export_to_excel


def staff(name, code):
    return dict(full_name=name, id=code, employee_code=code, active=True,
                chatwork_account_id='', aliases=[])


def directory():
    return {'initialized': True, 'employees': [
        staff('Lê Thu Bình', 'E-A'), staff('Nguyễn Hải Sơn', 'E-B'),
        staff('Trần Văn Sang', 'E-C'), staff('Lê Minh Quân', 'E-D'),
        staff('Võ Đức Quân', 'E-E'),
    ]}


def source(body, when='2026-10-09T08:00:00+07:00'):
    return dict(message_id='fake1', account_id='999', name='Người báo cáo (Demo)',
                send_time=when, body=body)


def records(body, when='2026-10-09T08:00:00+07:00'):
    msg = source(body, when)
    events = []
    reconcile_explicit_records(events, [msg], 10, 2026, roster_policy=True)
    for number, event in enumerate(events, 1):
        event['event_id'] = 'EVT' + str(number)
        classify(event, directory(), [msg])
    validate_set(events, [msg], 10, 2026)
    return msg, events


def test_contextual_minutes_belong_to_unique_roster_employee():
    _, events = records('Team Demo: ngày 2/10\nBình OT 2h30\nNam OT 2h30\nSơn 30ph đối ứng')
    by_name = {e['source_employee_name']: e for e in events}
    assert by_name['Sơn']['employee_name'] == 'Nguyễn Hải Sơn'
    assert by_name['Sơn']['ot_minutes'] == 30
    assert by_name['Sơn']['classification'] == 'official'
    assert not by_name['Sơn']['needs_review']
    assert not by_name['Sơn']['validation_errors']
    assert not by_name['Sơn']['validation_warnings']
    assert by_name['Nam']['classification'] == 'contractor'


def test_month_first_dates_in_october_ot_report():
    _, events = records('Team Demo đã tăng ca:\n10/03 Bình tăng ca 3h\n10/04 Bình tăng ca 3h',
                        '2026-10-04T11:00:00+07:00')
    assert {(e['event_date'], e['employee_name'], e['ot_minutes']) for e in events} == {
        ('2026-10-03', 'Lê Thu Bình', 180), ('2026-10-04', 'Lê Thu Bình', 180)}
    assert all(not e['needs_review'] and not e['validation_errors'] for e in events)


def test_prefix_group_short_names_do_not_need_full_name_quote():
    _, events = records('Team Demo: ngày 5/10\nBình OT 1h, anh Tâm + Sang OT 2h')
    by_name = {e['source_employee_name']: e for e in events}
    assert by_name['Sang']['employee_name'] == 'Trần Văn Sang'
    assert by_name['Sang']['classification'] == 'official'
    assert not by_name['Sang']['needs_review']


def test_namesakes_remain_review_even_after_parser_claim():
    _, events = records('Team Demo: ngày 8/10\nQuân OT 4h\nSơn OT 2h')
    by_name = {e['source_employee_name']: e for e in events}
    assert by_name['Quân']['classification'] == 'unresolved'
    assert by_name['Quân']['needs_review']
    assert by_name['Sơn']['classification'] == 'official'
    assert not by_name['Sơn']['needs_review']


def test_wrong_duration_cannot_be_verified_by_contextual_quote():
    msg, events = records('Team Demo: ngày 2/10\nBình OT 2h30\nSơn 30ph')
    son = next(e for e in events if e['source_employee_name'] == 'Sơn')
    son['ot_minutes'] = 60
    son.pop('employee_match_verified', None)
    son['needs_review'] = False
    son['review_reason'] = None
    classify(son, directory(), [msg])
    assert son['classification'] == 'official'
    assert son['needs_review']


def test_real_xlsx_has_official_and_contractor_not_only_source():
    msg, events = records('Team Demo: ngày 2/10\nBình OT 2h30\nNam OT 2h30\nSơn 30ph')
    stream = BytesIO()
    export_to_excel(events, {'month': 10, 'year': 2026, 'snapshot': [msg],
                             'ai_settings': {'employee_registry': directory()}}, stream)
    wb = openpyxl.load_workbook(BytesIO(stream.getvalue()), read_only=True, data_only=True)
    try:
        assert wb.sheetnames == ['1_TIN_NHAN_GOC','2_CHAM_CONG','3_TONG_KET',
                                 '4_CAN_KIEM_TRA','5_CTV']
        official = [row for row in wb['2_CHAM_CONG'].iter_rows(min_row=5, values_only=True)
                    if str(row[0] or '').startswith('EVT')]
        contractors = [row for row in wb['5_CTV'].iter_rows(min_row=5, values_only=True)
                       if str(row[0] or '').startswith('EVT')]
        assert {(row[5], row[19]) for row in official} == {
            ('Lê Thu Bình', 150), ('Nguyễn Hải Sơn', 30)}
        assert any(row[1] == 'Nam' and row[4] == 150 for row in contractors)
    finally:
        wb.close()
