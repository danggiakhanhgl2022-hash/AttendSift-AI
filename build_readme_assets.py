"""Build public README assets from synthetic data, never from a user's database.

Optional tooling: requirements-ui-test.txt + Pillow. Screenshots use installed
Edge. Generated workbook is held in memory; no employee workbook is committed.
"""
from pathlib import Path
from io import BytesIO
from html import escape
import argparse
import json
import logging
import sys
import tempfile
import textwrap
import threading

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / 'docs' / 'images' / 'story'


def fixtures():
    people = [dict(id=code, employee_code=code, full_name=name, aliases=[],
                   active=True, chatwork_account_id=account) for code, name, account in [
        ('DEMO01', 'Nguyễn Minh Quân', '101'),
        ('DEMO02', 'Trần Đức Quân', '102'),
        ('DEMO03', 'Lê Thu Bình', '103')]]
    def message(mid, name, body, day, account):
        return dict(message_id=mid, room_id='123', account_id=account, name=name,
                    body=body, send_time=f'2026-10-{day:02d}T16:30:00+07:00')
    messages = [
        message('DEMO-S01', 'Người báo cáo Demo',
                'Team Demo: ngày 14/10\nMinh Quân OT 2h\nĐức Quân OT 1h\nQuân OT 30p\nLan Demo (CTV) OT 45p', 15, '999'),
        message('DEMO-S02', 'Lê Thu Bình', 'Em xin phép về sớm 30 phút không làm bù.', 14, '103'),
        message('DEMO-S03', 'Lê Thu Bình', 'Em xin phép về sớm 30 phút không làm bù.', 16, '103'),
        message('DEMO-S04', 'Lê Thu Bình', 'Em xin phép nghỉ sáng ngày 15/10 (phép năm).', 14, '103'),
        message('DEMO-S05', 'Người báo cáo Demo', 'Nhắc cả nhóm kiểm tra lại báo cáo tháng.', 16, '999')]
    return {'initialized': True, 'revision': 1, 'employees': people}, messages


def svg(name, width, height, content):
    (OUT / name).write_text(
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}" role="img">'
        '<style>text{font-family:Segoe UI,Arial,sans-serif} .muted{fill:#9aafc0} .ink{fill:#edf5fa}</style>'
        + content + '</svg>', encoding='utf-8')


def diagrams():
    svg('cover.svg', 1440, 590, '''
    <defs><linearGradient id="bg"><stop stop-color="#09131f"/><stop offset="1" stop-color="#132f3d"/></linearGradient></defs>
    <rect width="1440" height="590" rx="24" fill="url(#bg)"/>
    <rect x="64" y="62" width="46" height="5" rx="2" fill="#55d9c1"/>
    <text x="64" y="105" fill="#55d9c1" font-size="17" letter-spacing="3">CHATWORK → ATTENDANCE</text>
    <text x="60" y="190" class="ink" font-size="76" font-weight="700">AttendSift AI</text>
    <text x="64" y="248" class="ink" font-size="29">Tin nhắn thành báo cáo.</text>
    <text x="64" y="290" class="ink" font-size="29">Mỗi kết quả đều có căn cứ.</text>
    <text x="64" y="352" class="muted" font-size="21">Windows · Cloud · Excel 5 sheet</text>
    <text x="64" y="520" class="muted" font-size="18">Đặng Gia Khánh / Bộ phận AI / 16.10.2026</text>
    <rect x="820" y="76" width="552" height="440" rx="18" fill="#0a1723" stroke="#335364"/>
    <text x="850" y="120" class="muted" font-size="16">MỘT VÍ DỤ GIẢ LẬP</text>
    <rect x="850" y="144" width="492" height="74" rx="10" fill="#193241"/>
    <text x="873" y="175" class="ink" font-size="20">Minh Quân OT 2h · Đức Quân OT 1h</text>
    <text x="873" y="202" class="muted" font-size="16">Ngày sự kiện: 14/10 · Source: DEMO-S01</text>
    <path d="M1096 229v30" stroke="#55d9c1" stroke-width="2"/>
    <text x="850" y="293" fill="#55d9c1" font-size="17">01  GIỮ NGUỒN</text>
    <text x="850" y="334" fill="#55d9c1" font-size="17">02  ĐỐI CHIẾU DANH SÁCH</text>
    <text x="850" y="375" fill="#55d9c1" font-size="17">03  KIỂM TRA QUY TẮC</text>
    <line x1="850" y1="399" x2="1340" y2="399" stroke="#335364"/>
    <text x="850" y="439" class="ink" font-size="25" font-weight="600">120 phút + 60 phút</text>
    <text x="850" y="474" class="muted" font-size="17">Tách từng người. Không chia OT của cả nhóm.</text>''')
    blocks = [('01', 'Tin nguồn', 'Giữ nội dung và Source ID'),
              ('02', 'Đối chiếu', 'Người · ngày · thời lượng'),
              ('03', 'Phân nhóm', 'Chính thức · CTV · cần kiểm tra'),
              ('04', 'Excel', '5 sheet để đọc và đối chiếu')]
    parts = ['<rect width="1440" height="350" rx="22" fill="#f1f6fa"/>',
             '<text x="40" y="60" fill="#142d40" font-size="29" font-weight="700">Kiến trúc bắt đầu từ câu hỏi của con người</text>']
    for i, (number, title, subtitle) in enumerate(blocks):
        x = 40 + i * 350
        parts.append(f'<rect x="{x}" y="104" width="310" height="162" rx="14" fill="white" stroke="#cadce7"/>'
                     f'<text x="{x+22}" y="139" fill="#137f75" font-size="18">{number}</text>'
                     f'<text x="{x+22}" y="183" fill="#142d40" font-size="27" font-weight="600">{title}</text>'
                     f'<text x="{x+22}" y="225" fill="#4b6476" font-size="16">{escape(subtitle)}</text>')
        if i < 3:
            parts.append(f'<path d="M{x+318} 185h23m-8-7 8 7-8 7" fill="none" stroke="#137f75" stroke-width="2"/>')
    parts.append('<text x="40" y="311" fill="#4b6476" font-size="19">AI gợi ý sự kiện. Quy tắc kiểm tra. Người phụ trách xác nhận trường hợp chưa rõ.</text>')
    svg('journey.svg', 1440, 350, ''.join(parts))


def workbook_images():
    from PIL import Image, ImageDraw, ImageFont
    import openpyxl
    from cloud.app import finalize
    from src.utils.excel_export import export_to_excel
    roster, messages = fixtures()
    events = finalize([], messages, 10, 2026, 'DEMO-PUBLIC', employee_registry=roster)
    official = [e for e in events if e.get('classification') == 'official' and not e.get('needs_review')]
    assert sum(e.get('ot_minutes', 0) for e in official) == 180
    assert sum(e.get('paid_break_minutes', 0) for e in official) == 60
    assert len([e for e in events if e.get('classification') == 'unresolved']) == 1
    stream = BytesIO()
    export_to_excel(events, {'month': 10, 'year': 2026, 'snapshot': messages,
        'session_id': 'DEMO-PUBLIC', 'is_draft': True,
        'ai_settings': {'employee_registry': roster}}, stream)
    wb = openpyxl.load_workbook(BytesIO(stream.getvalue()), data_only=True)
    font_dir = Path('C:/Windows/Fonts')
    def font(size, bold=False):
        if font_dir.exists():
            return ImageFont.truetype(str(font_dir / ('arialbd.ttf' if bold else 'arial.ttf')), size)
        return ImageFont.truetype('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf', size)
    specs = [
        ('01-source.png', '1_TIN_NHAN_GOC', '01 / Tin nhắn gốc',
         'Ngày gửi giữ riêng với ngày sự kiện; tin báo hộ không biến thành công của người gửi.',
         [1, 4, 5, 7], [190, 260, 750, 300]),
        ('02-attendance.png', '2_CHAM_CONG', '02 / Chi tiết chính thức',
         'Cột được chọn từ workbook; nghỉ có lương tách khỏi phút thiếu giờ và OT.',
         [2, 5, 6, 7, 15, 20, 28], [185, 180, 285, 180, 210, 210, 250]),
        ('03-summary.png', '3_TONG_KET', '03 / Tổng kết',
         'Chỉ cộng sự kiện đủ căn cứ: 180 phút OT chính thức và 60 phút nghỉ có lương.',
         [1, 8, 9, 12, 15], [360, 240, 300, 300, 300]),
        ('04-review.png', '4_CAN_KIEM_TRA', '04 / Cần kiểm tra',
         'Tên “Quân” khớp hai nhân viên: giữ căn cứ, chờ chọn đúng người; chưa cộng công.',
         [1, 3, 4, 5, 6], [190, 190, 220, 650, 350]),
        ('05-contractors.png', '5_CTV', '05 / Cộng tác viên',
         'OT CTV 45 phút được giữ riêng, không cộng vào tổng OT chính thức.',
         [2, 3, 4, 5, 7], [290, 230, 230, 220, 630])]
    for filename, sheetname, title, subtitle, columns, widths in specs:
        sheet = wb[sheetname]
        width, height = sum(widths) + 80, 830
        image = Image.new('RGB', (width, height), '#f2f6f9')
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, width, 160), fill='#102a3d')
        draw.text((40, 26), title, font=font(34, True), fill='white')
        draw.text((40, 87), subtitle, font=font(21), fill='#c6e3ee')
        draw.text((40, 170), 'DỮ LIỆU GIẢ LẬP · Trình bày rút gọn từ workbook do mã nguồn xuất', font=font(18, True), fill='#137f75')
        x = 40
        for col, cell_width in zip(columns, widths):
            draw.rectangle((x, 212, x + cell_width, 296), fill='#205277')
            header = str(sheet.cell(4, col).value or '')
            for n, line in enumerate(textwrap.wrap(header, max(12, cell_width // 14))[:3]):
                draw.text((x + 12, 226 + n * 22), line, font=font(19, True), fill='white')
            x += cell_width
        y = 296
        for row in range(5, min(sheet.max_row + 1, 10)):
            if sheetname == '5_CTV' and sheet.cell(row, 1).value is None:
                break
            if not any(sheet.cell(row, col).value is not None for col in columns):
                continue
            rowheight = 82 if sheetname != '1_TIN_NHAN_GOC' else 90
            x = 40
            for col, cell_width in zip(columns, widths):
                draw.rectangle((x, y, x + cell_width, y + rowheight), fill='#d9eff6' if row % 2 else 'white', outline='#c9dce6')
                value = sheet.cell(row, col).value
                if hasattr(value, 'strftime'):
                    value = value.strftime('%d/%m/%Y')
                raw = str(value if value is not None else '')
                lines = []
                for line in raw.splitlines():
                    lines += textwrap.wrap(line, max(12, cell_width // 12)) or ['']
                shown = lines[:3]
                if len(lines) > 3:
                    shown[-1] = shown[-1].rstrip() + '…'
                for n, line in enumerate(shown):
                    draw.text((x + 12, y + 11 + n * 22), line, font=font(19), fill='#183446')
                x += cell_width
            y += rowheight
        draw.text((40, height - 76), 'Các ảnh không chứa tên nhân viên, khách hàng hoặc tin nhắn thật.', font=font(19), fill='#4b6476')
        draw.text((40, height - 42), 'Minh họa được dựng từ giá trị workbook; không phải ảnh chụp Microsoft Excel hoặc WPS.', font=font(17), fill='#4b6476')
        image.save(OUT / filename, optimize=True)
    wb.close()
    (OUT / 'demo-manifest.json').write_text(json.dumps({
        'synthetic': True, 'source_count': len(messages), 'official_ot_minutes': 180,
        'official_paid_break_minutes': 60, 'contractor_ot_minutes': 45,
        'unresolved_events': 1, 'notes': 'No live API calls; no production database.'}, indent=2), encoding='utf-8')


def screenshots():
    from playwright.sync_api import sync_playwright
    from werkzeug.serving import make_server
    from src.app.server import create_app
    from unittest.mock import Mock, patch
    roster, messages = fixtures()
    logging.getLogger('werkzeug').setLevel(logging.ERROR)
    with tempfile.TemporaryDirectory(prefix='attendsift-docs-') as temp:
        app = create_app(temp)
        services = app.extensions['services']
        services['config'].set('chatwork_room_id', '123')
        services['config'].set('chatwork_room_name', 'Nhóm Demo — dữ liệu giả lập')
        services['employees'].import_entries([{k: v for k, v in person.items() if k != 'id'}
                                               for person in roster['employees']])
        services['db'].now_vn = lambda: '2026-10-16T16:00:00+07:00'
        from src.models.database import MessageRepository
        MessageRepository(services['db']).save_messages_batch(messages)
        export = services['export']
        # Checkpoints and history use synthetic extraction, never a remote provider.
        with patch.object(export.config_repo, 'get_ai_provider', return_value={
                'provider': 'custom', 'model': 'demo-offline', 'api_key': 'synthetic-offline-only', 'base_url': ''}):
            sid = export.create_session(10, 2026, True)['session_id']
            adapter = Mock()
            adapter.analyze_attendance.return_value = []
            with patch('src.services.export_service.create_ai_adapter', return_value=adapter):
                export.analyze(sid)
        server = make_server('127.0.0.1', 0, app, threaded=True)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        errors = []
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(channel='msedge', headless=True)
                context = browser.new_context(viewport={'width': 1600, 'height': 1000}, locale='vi-VN', device_scale_factor=1)
                context.route('**/*', lambda route: route.continue_() if
                    route.request.url.startswith(('http://127.0.0.1:', 'data:', 'blob:')) else route.abort())
                page = context.new_page()
                page.on('pageerror', lambda error: errors.append(str(error)))
                page.goto(f'http://127.0.0.1:{server.server_port}', wait_until='networkidle')
                page.locator('.nav-item').nth(4).wait_for(state='visible')
                page.locator('.nav-item[data-section="sync"]').click()
                page.locator('#sync-section').wait_for(state='visible')
                page.wait_for_timeout(1500)
                for section in ['sync', 'export', 'history', 'settings', 'employees']:
                    page.locator(f'.nav-item[data-section="{section}"]').click()
                    page.locator(f'#{section}-section').wait_for(state='visible')
                    page.wait_for_timeout(700)
                    if section == 'export':
                        page.locator('#month').fill('2026-10')
                    page.screenshot(path=str(OUT / f'ui-{section}.png'))
                assert not errors, errors
                context.close()
                browser.close()
        finally:
            server.shutdown()
            thread.join(timeout=5)
            server.server_close()
    return errors


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--screenshots', action='store_true')
    args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    diagrams()
    workbook_images()
    if args.screenshots:
        errors = screenshots()
        (OUT / 'ui-capture-report.json').write_text(json.dumps({
            'browser': 'Microsoft Edge headless', 'synthetic_database': True,
            'tabs_captured': ['sync', 'export', 'history', 'settings', 'employees'],
            'javascript_errors': errors, 'live_api_called': False,
            'note': 'Browser screenshots do not certify WebView2 or Cloud deployment.'
        }, indent=2), encoding='utf-8')
    print('Public assets generated from synthetic fixtures:', OUT)
